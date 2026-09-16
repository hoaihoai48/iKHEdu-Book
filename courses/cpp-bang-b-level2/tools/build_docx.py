#!/usr/bin/env python3
"""
Build script: Tạo file Word (.docx) xuất bản từ MASTER_ALL_LESSONS.md
Khoá học C++ Bảng B (Level 2) — Nâng Cao Tư Duy Thuật Toán & Kỹ Thuật Thi Đấu
Trung tâm tin học iKH

Hỗ trợ xuất bản theo 2 Quyển độc lập chuẩn in ấn:
- QUYỂN 1: Chương 01 -> 03 (Bài 01 -> 06) + Phụ lục A + Bài giải 01–06 (134 bài toán thực hành)
- QUYỂN 2: Chương 04 -> 06 (Bài 07 -> 15) + Phụ lục A + Bài giải 07–15 (212 bài toán thực hành)
TỔNG CỘNG: 346 bài toán thực hành phân tầng chất lượng cao kèm 100% mã giải thuật C++ chuẩn thi đấu.
"""

import os
import re
import sys
import glob
import html
import argparse
import subprocess
from pathlib import Path

# ============================================================
# CONFIGURATION
# ============================================================
BOOK_AUTHOR = "Trung tâm tin học iKH"

BASE_DIR = Path(__file__).parent
MASTER_FILE = BASE_DIR / "MASTER_ALL_LESSONS.md"
FOUNDATION_FILE = BASE_DIR.parent / "cpp-bang-b" / "source" / "level0" / "IKHEDU_Level0_Foundation.md"
PROBLEMS_DIR = BASE_DIR / "problems"

ALL_TOPICS = [
    # Quyển 1: Số Học Nâng Cao, Kỹ Thuật Mảng & Bitwise (Chương 01 - 03, Bài 01 - 06: 134 bài)
    ("Chương 01 — Bài 01: Số học cơ bản & chuyên sâu", "cppb2_l01", 1),
    ("Chương 01 — Bài 02: Modulo & lũy thừa nhanh", "cppb2_l02", 1),
    ("Chương 02 — Bài 03: Tìm kiếm nhị phân nâng cao", "cppb2_l03", 1),
    ("Chương 02 — Bài 04: Kỹ thuật mảng: Two Pointers, Window & 2D Prefix", "cppb2_l04", 1),
    ("Chương 03 — Bài 05: Đệ quy, chia để trị & Meet in the Middle", "cppb2_l05", 1),
    ("Chương 03 — Bài 06: Phép toán bit & mặt nạ bit nâng cao", "cppb2_l06", 1),

    # Quyển 2: Thuật Toán Tối Ưu, Cấu Trúc Dữ Liệu & Đồ Thị (Chương 04 - 06, Bài 07 - 15: 212 bài)
    ("Chương 04 — Bài 07: Thuật toán tham lam (Greedy)", "cppb2_l07", 2),
    ("Chương 04 — Bài 08: Quy hoạch động cơ bản & chuyên sâu (DP)", "cppb2_l08", 2),
    ("Chương 05 — Bài 09: Ngăn xếp, hàng đợi & Deque đơn điệu", "cppb2_l09", 2),
    ("Chương 05 — Bài 10: Thư viện STL C++ nâng cao", "cppb2_l10", 2),
    ("Chương 05 — Bài 11: Tổ hợp, hoán vị & xác suất cơ bản", "cppb2_l11", 2),
    ("Chương 06 — Bài 12: Lý thuyết đồ thị cơ bản & nâng cao", "cppb2_l12", 2),
    ("Chương 06 — Bài 13: Cây phân đoạn & Cây Fenwick", "cppb2_l13", 2),
    ("Chương 06 — Bài 14: Quy hoạch động chữ số (Digit DP)", "cppb2_l14", 2),
    ("Chương 06 — Bài 15: Xử lý chuỗi, String Hashing & BigInt", "cppb2_l15", 2),
]

CODE_TO_DIR_PREFIX = {f"CPPB2-L{i:02d}": f"cppb2_l{i:02d}" for i in range(1, 16)}


# ============================================================
# PHASE 0: SVG & Image Processing
# ============================================================

def convert_svg_to_png():
    svg_files = list(BASE_DIR.glob("lessons/**/assets/*.svg"))
    converted_count = 0

    for svg in svg_files:
        png = svg.with_suffix(".png")
        if not png.exists() or png.stat().st_mtime < svg.stat().st_mtime:
            cmd = [
                "npx", "-y", "@resvg/resvg-js-cli",
                "--fit-width", "2800",
                "--dpi", "300",
                "--text-rendering", "2",
                "--shape-rendering", "2",
                str(svg),
                str(png)
            ]
            res = subprocess.run(cmd, capture_output=True)
            if res.returncode == 0 and png.exists():
                converted_count += 1

    if converted_count > 0:
        print(f"  → Đã chuyển đổi {converted_count} hình ảnh SVG sang PNG siêu nét 2800px (300 DPI)")


def resolve_and_embed_images(content):
    png_map = {}
    for p in BASE_DIR.glob("lessons/**/assets/*.png"):
        png_map[p.name] = str(p)
        svg_name = p.stem + ".svg"
        png_map[svg_name] = str(p)

    def replace_img(match):
        alt = match.group(1)
        orig_path = match.group(2)
        filename = Path(orig_path).name

        if filename in png_map:
            return f"\n\n![{alt}]({png_map[filename]})\n\n"
        return match.group(0)

    content = re.sub(r"!\[(.*?)\]\((.*?)\)", replace_img, content)
    return content


# ============================================================
# PHASE 1: Pre-process markdown
# ============================================================

def read_file(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()

def find_problem_dir(code):
    match = re.match(r"(CPPB2-L\d+)-(\d+)", code)
    if not match:
        code_pattern = code.lower().replace("-", "_")
        matches = list(PROBLEMS_DIR.glob(f"{code_pattern}_*"))
        return matches[0] if matches else None
    
    prefix = match.group(1)
    num = match.group(2)
    dir_prefix = CODE_TO_DIR_PREFIX.get(prefix)
    if not dir_prefix:
        return None
    
    pattern = f"{dir_prefix}_{num}_*"
    matches = list(PROBLEMS_DIR.glob(pattern))
    if matches:
        return matches[0]
    return None

def read_problem_statement(problem_dir):
    de_bai = problem_dir / "De_Bai.md"
    if de_bai.exists():
        return read_file(de_bai)
    return None

def read_solution(problem_dir):
    sol = problem_dir / "solution.cpp"
    if sol.exists():
        return read_file(sol)
    return None

def extract_problem_codes_from_table(table_text):
    return re.findall(r"(CPPB2-L\d+-\d+)", table_text)

def remove_quiz_sections(lines):
    result = []
    skip = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if re.match(r"^##\s+.*(?:Câu hỏi trắc nghiệm|Hệ thống câu hỏi)", line, re.IGNORECASE):
            skip = True
            i += 1
            continue
        if skip:
            if re.match(r"^#{1,2}\s", line) and not re.match(r"^###", line):
                skip = False
                result.append(line)
            i += 1
            continue
        result.append(line)
        i += 1
    return result

def remove_overview_toc(lines):
    result = []
    skip = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip() == "## MỤC LỤC TỔNG QUAN":
            skip = True
            i += 1
            continue
        if skip:
            if line.startswith("===") or line.startswith("# MODULE") or line.startswith("# CHƯƠNG") or "BẮT ĐẦU NỘI DUNG" in line:
                skip = False
                result.append(line)
            i += 1
            continue
        result.append(line)
        i += 1
    return result

def remove_top_header_and_metadata(lines):
    result = []
    i = 0
    skip_header = True
    while i < len(lines):
        line = lines[i]
        if skip_header:
            if line.startswith("# MODULE") or line.startswith("# CHƯƠNG") or "BẮT ĐẦU NỘI DUNG" in line:
                skip_header = False
                if "BẮT ĐẦU NỘI DUNG" in line:
                    i += 1
                    continue
                result.append(line)
            i += 1
            continue
        result.append(line)
        i += 1
    return result

def _split_sample_blocks(v):
    """Tách khối Sample thành Input / Output / Giải thích theo các dòng ###."""
    blocks = {"input": [], "output": [], "giaithich": [], "other": []}
    cur = "other"
    for line in v.split("\n"):
        s = line.strip()
        if re.match(r"^#{1,4}\s*Input\b", s, re.IGNORECASE):
            cur = "input"
            continue
        if re.match(r"^#{1,4}\s*Output\b", s, re.IGNORECASE):
            cur = "output"
            continue
        if re.match(r"^#{1,4}\s*Giải thích\b", s, re.IGNORECASE):
            cur = "giaithich"
            continue
        if re.match(r"^#{1,4}\s*\S", s):
            cur = "other"
        blocks[cur].append(line)
    return blocks

def _md_table_cell(lines):
    cleaned = [l.strip().replace("|", "\\|") for l in lines if l.strip() != ""]
    return "<br/>".join(cleaned) if cleaned else ""

def format_problem_statement(statement, code, idx=1):
    lines = statement.strip().split("\n")
    title = ""
    sections = {}
    current_sec = "other"
    sec_content = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            title = line[2:].strip()
            title = re.sub(r"##\s*Mã bài toán.*", "", title).strip()
            i += 1
            continue
        m_h2 = re.match(r"^##\s+(?:📖\s*|📥\s*|📤\s*|📌\s*|⚙️\s*)?(?:\d+\.\s*)?(.+)$", line)
        if m_h2:
            if sec_content:
                sections[current_sec] = "\n".join(sec_content).strip()
            current_sec = m_h2.group(1).strip()
            sec_content = []
            i += 1
            continue
        sec_content.append(line)
        i += 1
        
    if sec_content:
        sections[current_sec] = "\n".join(sec_content).strip()
        
    out = []
    clean_title = title.split("\n")[0].strip()
    out.append(f"### Bài {idx:02d} [{code}]: {clean_title}\n")
    
    for k, v in sections.items():
        if "Bối cảnh" in k or "Nhiệm vụ" in k:
            out.append(f"**Bối cảnh & Nhiệm vụ:**\n\n{v}\n")
        elif "Đầu Vào" in k or "Input" in k:
            out.append(f"**Đầu vào (Input):**\n\n{v}\n")
        elif "Đầu Ra" in k or "Output" in k:
            out.append(f"**Đầu ra (Output):**\n\n{v}\n")
        elif "Ví Dụ" in k or "Sample" in k:
            sb = _split_sample_blocks(v)
            cell_in = _md_table_cell(sb["input"])
            cell_out = _md_table_cell(sb["output"])
            out.append("**Ví dụ mẫu:**\n\n| Đầu vào (Input) | Đầu ra (Output) |\n|---|---|\n"
                       f"| {cell_in} | {cell_out} |\n")
            gt_lines = [l.strip() for l in sb["giaithich"] if l.strip() != ""]
            if gt_lines:
                out.append("**Giải thích:**\n\n" + "\n\n".join(gt_lines) + "\n")
            other_lines = [l for l in sb["other"] if l.strip() != ""]
            if other_lines:
                out.append("\n".join(other_lines) + "\n")
        elif "Ràng Buộc" in k or "Constraints" in k:
            clean_rb = "\n".join([l for l in v.split("\n") if not re.search(r"Thời gian|Bộ nhớ", l, re.IGNORECASE)])
            if clean_rb.strip():
                out.append(f"**Ràng buộc dữ liệu:**\n\n{clean_rb.strip()}\n")
                
    return "\n".join(out) + "\n\n"

def replace_exercise_matrices_with_problems(lines):
    result = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if re.match(r"^##\s+.*Ma trận.*bài tập", line, re.IGNORECASE):
            section_lines = [line]
            i += 1
            while i < len(lines):
                next_line = lines[i]
                if (next_line.startswith("# ") or next_line.startswith("# CHƯƠNG") or 
                    next_line.startswith("# Bài") or next_line.startswith("<!--") or 
                    re.match(r"^={3,}$", next_line.strip())):
                    break
                section_lines.append(next_line)
                i += 1
            section_text = "\n".join(section_lines)
            codes = extract_problem_codes_from_table(section_text)
            
            result.append("## Bài tập thực hành phân tầng (P0 → P5)\n\n")
            for idx, code in enumerate(codes, 1):
                pdir = find_problem_dir(code)
                if pdir:
                    stmt = read_problem_statement(pdir)
                    if stmt:
                        result.append(format_problem_statement(stmt, code, idx))
                    else:
                        result.append(f"### Bài {idx:02d} [{code}]\n\n*Đề bài chưa có.*\n\n")
                else:
                    result.append(f"### Bài {idx:02d} [{code}]\n\n*Đề bài chưa có.*\n\n")
            continue
        result.append(line)
        i += 1
    return result

def ensure_markdown_list_blank_lines(lines):
    result = []
    list_pattern = re.compile(r"^(?:[\*\-\+]|\d+\.)\s+")
    bq_list_pattern = re.compile(r"^>\s*(?:[\*\-\+]|\d+\.)\s+")
    for i, line in enumerate(lines):
        stripped = line.strip()
        is_list = bool(list_pattern.match(stripped))
        is_bq_list = bool(bq_list_pattern.match(stripped))
        if is_list and len(result) > 0:
            prev = result[-1].strip()
            prev_is_list = bool(list_pattern.match(prev))
            if prev and not prev_is_list and not prev.startswith("```"):
                result.append("")
        elif is_bq_list and len(result) > 0:
            prev = result[-1].strip()
            prev_is_bq_list = bool(bq_list_pattern.match(prev))
            if prev and not prev_is_bq_list and not prev.startswith("```") and not prev.startswith("> ```"):
                result.append(">")
        result.append(line)
    return result

def clean_separators(lines):
    result = []
    for line in lines:
        stripped = line.strip()
        if re.match(r"^={3,}$", stripped) or re.match(r"^-{3,}$", stripped):
            continue
        if re.match(r"^<!--.*-->$", stripped):
            continue
        result.append(line)
    return result

def format_callouts_and_blockquotes(content):
    lines = content.split("\n")
    new_lines = []
    for i, line in enumerate(lines):
        new_lines.append(line)
    return "\n".join(new_lines)

def build_front_matter(volume=1):
    if volume == 1:
        vol_name = "QUYỂN 1: SỐ HỌC NÂNG CAO, KỸ THUẬT MẢNG & BITWISE"
        vol_desc = "gồm **3 Chương trọng tâm (Chương 01 đến Chương 03)** với **6 Bài học** và **134 bài toán thực hành phân tầng (P0 → P5)**, đào sâu số học đồng dư, lũy thừa nhanh, tìm kiếm nhị phân không gian nghiệm, kỹ thuật mảng 2D, đệ quy chia để trị, Meet in the Middle và mặt nạ bit."
    elif volume == 2:
        vol_name = "QUYỂN 2: THUẬT TOÁN TỐI ƯU, CẤU TRÚC DỮ LIỆU & ĐỒ THỊ"
        vol_desc = "gồm **3 Chương chuyên sâu (Chương 04 đến Chương 06)** với **9 Bài học** và **212 bài toán thực hành phân tầng (P0 → P5)**, chinh phục tham lam, quy hoạch động tối ưu, ngăn xếp đơn điệu, Deque, STL C++ nâng cao, tổ hợp, đồ thị BFS/DFS/Dijkstra, cây phân đoạn Segment Tree, Fenwick Tree, Digit DP và xử lý chuỗi Hashing."
    else:
        vol_name = "TRỌN BỘ 6 CHƯƠNG"
        vol_desc = "bao gồm đầy đủ 6 Chương chuyên sâu với 15 Bài học và 346 bài toán thực hành có lời giải chi tiết."

    return f"""\\newpage

# Lời nói đầu

Chào mừng các em học sinh và quý thầy cô đến với bộ giáo trình **Khoá học C++ Bảng B (Level 2) — {vol_name}** của Trung tâm tin học iKH.

Bộ tài liệu này được biên soạn công phu nhằm cung cấp lộ trình học tập lập trình thi đấu nâng cao, chuẩn mực và hiện đại nhất dành cho học sinh giỏi Tin học THCS, THPT và sinh viên Olympic Tin học.

Phần nội dung này {vol_desc}

Mỗi bài học được thiết kế theo cấu trúc sư phạm chặt chẽ:

- **Khái niệm & Bản chất toán học**: Giải thích trực quan, dễ hiểu kèm chứng minh toán học và bất biến thuật toán.
- **Bảng mô phỏng từng bước (Dry Run Table)**: Trực quan hóa quá trình biến đổi dữ liệu từng bước.
- **Mẫu cài đặt chuẩn thi đấu**: Code C++ chuẩn (0 `std::`, `#include <bits/stdc++.h>`, Fast I/O), tối ưu và an toàn tuyệt đối.
- **Hệ thống bài tập thực hành phân tầng**: Từ cơ bản đến chuyên sâu (P0 đến P5), có đầy đủ bối cảnh, nhiệm vụ, input/output và sample test.
- **Lời giải tham khảo chi tiết**: Phụ lục B cung cấp toàn bộ mã nguồn C++ tham khảo chuẩn thi đấu cho các bài tập.

Chúc các em học tập hiệu quả và chinh phục những giải thưởng cao trong các kỳ thi chọn học sinh giỏi và Olympic lập trình!

\\begin{{flushright}}
\\textbf{{{BOOK_AUTHOR}}}
\\end{{flushright}}
"""

def build_back_matter_foundation():
    if not FOUNDATION_FILE.exists():
        return ""
    content = read_file(FOUNDATION_FILE)
    result = "\n\\newpage\n\n"
    result += "# Phụ lục A: Nền tảng C++\n\n"
    result += "> Phần này tóm tắt toàn bộ cú pháp, cấu trúc dữ liệu và quy trình giải bài C++ cơ bản.\n\n"
    result += content
    return result

def collect_solutions_for_volume(volume=None):
    if volume == 1:
        target_topics = [t for t in ALL_TOPICS if t[2] == 1]
    elif volume == 2:
        target_topics = [t for t in ALL_TOPICS if t[2] == 2]
    else:
        target_topics = ALL_TOPICS

    result = "\n\\newpage\n\n"
    result += "# Phụ lục B: Lời giải bài tập tham khảo\n\n"
    result += "> Phần này cung cấp mã nguồn C++ tham khảo chuẩn thi đấu cho các bài tập thực hành trong sách.\n\n"

    total_solutions = 0
    for topic_title, dir_prefix, vol in target_topics:
        topic_pattern = f"{dir_prefix}_*"
        dirs = sorted(PROBLEMS_DIR.glob(topic_pattern))
        if not dirs:
            continue

        result += f"## {topic_title}\n\n"
        for pdir in dirs:
            sol_file = pdir / "solution.cpp"
            if not sol_file.exists():
                continue
            de_bai_file = pdir / "De_Bai.md"
            problem_name = pdir.name
            if de_bai_file.exists():
                first_line = read_file(de_bai_file).split("\n")[0]
                problem_name = first_line.replace("#", "").strip()

            parts = pdir.name.split("_")
            code = parts[0].upper() + "-" + parts[1].upper() + "-" + parts[2]
            solution_code = read_file(sol_file)

            result += f"### `{code}` — {problem_name}\n\n"
            result += f"```cpp\n{solution_code}\n```\n\n"
            total_solutions += 1

    print(f"  → Đã thu thập {total_solutions} bài giải cho {'Quyển ' + str(volume) if volume else 'Toàn bộ'}")
    return result

def extract_volume_lines(lines, volume=1):
    if volume == 1:
        vol_lines = []
        for line in lines:
            if line.startswith("# CHƯƠNG 04"):
                break
            vol_lines.append(line)
        return vol_lines
    elif volume == 2:
        vol_lines = []
        capture = False
        for line in lines:
            if line.startswith("# CHƯƠNG 04"):
                capture = True
            if capture:
                vol_lines.append(line)
        return vol_lines
    else:
        return lines

def preprocess_markdown_for_volume(volume=1):
    vol_str = f"QUYỂN {volume}" if volume in [1, 2] else "TRỌN BỘ"
    print(f"\n📖 Đang xử lý nội dung cho {vol_str}...")

    convert_svg_to_png()
    content = read_file(MASTER_FILE)
    content = re.sub(r"# MODULE (\d+)", r"# CHƯƠNG \1", content)
    content = re.sub(r"MODULE (\d+)", r"CHƯƠNG \1", content)
    content = re.sub(r"# Chuyên đề (\d+)", r"# Bài \1", content)
    content = re.sub(r"Chuyên đề (\d+)", r"Bài \1", content)

    content = re.sub(r"#\s+include\b", "#include", content)
    content = re.sub(r"#\s+define\b", "#define", content)
    content = re.sub(r"#\s+pragma\b", "#pragma", content)
    content = re.sub(r"#\s+ifndef\b", "#ifndef", content)
    content = re.sub(r"#\s+endif\b", "#endif", content)
    content = re.sub(r"#\s+ifdef\b", "#ifdef", content)

    content = format_callouts_and_blockquotes(content)
    content = resolve_and_embed_images(content)

    lines = content.split("\n")
    lines = remove_overview_toc(lines)
    lines = remove_top_header_and_metadata(lines)
    lines = remove_quiz_sections(lines)
    lines = replace_exercise_matrices_with_problems(lines)
    lines = clean_separators(lines)
    lines = ensure_markdown_list_blank_lines(lines)

    lines = extract_volume_lines(lines, volume)
    print(f"  → Nội dung {vol_str}: {len(lines)} dòng")

    front = build_front_matter(volume)
    foundation = build_back_matter_foundation() if volume in [1, 2, None] else ""
    solutions = collect_solutions_for_volume(volume)
    toc_section = "\n\\newpage\n\n# Mục lục\n\n"

    body = "\n".join(lines)

    if volume == 1:
        book_title = "Khoá học C++ Bảng B (Level 2) — Quyển 1"
        book_subtitle = "Số học nâng cao, Kỹ thuật mảng & Bitwise (Bài 01-06)"
    elif volume == 2:
        book_title = "Khoá học C++ Bảng B (Level 2) — Quyển 2"
        book_subtitle = "Thuật toán tối ưu, Cấu trúc dữ liệu & Đồ thị (Bài 07-15)"
    else:
        book_title = "Khoá học C++ Bảng B (Level 2)"
        book_subtitle = "Nâng cao tư duy, Cấu trúc dữ liệu & Kỹ thuật thi đấu"

    # B2.1: KHÔNG sinh trang bìa — bỏ title/subtitle/author khỏi YAML
    # để pandoc không phát sinh title block (sách mở đầu thẳng vào Lời nói đầu).
    final_md = f"""---
lang: vi
documentclass: report
geometry: "a4paper, margin=2.5cm"
fontsize: 12pt
mainfont: "Times New Roman"
monofont: "Courier New"
header-includes:
  - \\usepackage{{fancyhdr}}
  - \\pagestyle{{fancy}}
  - \\fancyhead[L]{{\\textit{{{book_title}}}}}
  - \\fancyhead[R]{{\\textit{{{BOOK_AUTHOR}}}}}
---

{front}

{body}

{foundation}

{solutions}

{toc_section}
"""
    return final_md, book_title, book_subtitle

def build_with_pandoc(markdown_content, output_file):
    print(f"\n🔧 Đang chạy Pandoc chuyển đổi → {output_file.name}...")
    temp_md = BASE_DIR / f"_build_intermediate_{output_file.stem}.md"
    with open(temp_md, "w", encoding="utf-8") as f:
        f.write(markdown_content)

    print(f"  → Đã ghi file trung gian: {temp_md.name} ({temp_md.stat().st_size:,} bytes)")

    cmd = [
        "pandoc",
        str(temp_md),
        "-o", str(output_file),
        "--from=markdown+tex_math_dollars+pipe_tables+fenced_code_blocks+yaml_metadata_block+header_attributes",
        "--to=docx",
        "--standalone",
        "--wrap=auto",
        "--columns=80",
    ]

    ref_doc = BASE_DIR.parent / "cpp-bang-b" / "_reference.docx"
    if ref_doc.exists():
        cmd.extend(["--reference-doc", str(ref_doc)])

    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        print(f"  ❌ Pandoc lỗi: {result.stderr}")
        return False

    print(f"  ✅ Pandoc thành công!")
    return True

WATERMARK_LOGO = BASE_DIR / "watermark_logo.jpeg"  # tách từ image13.jpeg Quyển 1 chuẩn

def _add_vml_watermark(header, logo_path):
    """B3.1: Chèn Watermark VML (logo mờ 286pt, căn giữa trang, nằm dưới chữ)."""
    from docx.shared import Pt
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    p = header.add_paragraph()
    run = p.add_run()
    run.add_picture(str(logo_path), width=Pt(286))
    drawing = run._r.find(qn('w:drawing'))
    if drawing is None:
        return
    blip = drawing.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    if blip is None:
        return
    rId = blip.get(qn('r:embed'))
    vml_ns = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:w10="urn:schemas-microsoft-com:office:word" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    )
    pict = parse_xml(
        f'<w:pict {vml_ns}>'
        '<v:shape id="WordPictureWatermark" type="#_x0000_t75" '
        'style="position:absolute;margin-left:0;margin-top:0;width:286pt;height:286pt;'
        'z-index:-251656192;mso-position-horizontal:center;'
        'mso-position-horizontal-relative:margin;mso-position-vertical:center;'
        'mso-position-vertical-relative:margin" o:allowincell="f">'
        f'<v:imagedata r:id="{rId}" o:title="logo_ikh" gain="19661f" blacklevel="22938f"/>'
        '<w10:wrap anchorx="margin" anchory="margin"/>'
        '</v:shape></w:pict>')
    drawing.addprevious(pict)
    drawing.getparent().remove(drawing)

def postprocess_docx(output_file, book_title, book_subtitle, volume=1):
    print(f"\n🎨 Đang áp dụng Design System xuất bản chuyên nghiệp cho {output_file.name}...")

    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml, OxmlElement
    from docx.oxml.ns import nsdecls, qn

    doc = Document(str(output_file))
    # B3.1: bật evenAndOddHeaders để Word tôn trọng header trang chẵn
    doc.settings.odd_and_even_pages_header_footer = True

    # ============================================================
    # 1. PAGE SETUP & MARGINS (Đồng bộ chuẩn 100% Level 1)
    # ============================================================
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        # B2.4: Margins chuẩn in — Top/Bottom/Right 36pt, Left (gáy) 64.35pt
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
        section.left_margin = Pt(64.35)
        section.right_margin = Pt(36)
        section.different_first_page_header_footer = True

        # Header text cho default + even (Spec §2.2: trái iKHEDU / phải tên sách, 9.5pt italic)
        for hdr in (section.header, section.even_page_header):
            p_head = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
            p_head.text = ""
            p_head.paragraph_format.tab_stops.add_tab_stop(Cm(17.46), WD_TAB_ALIGNMENT.RIGHT)
            r_hleft = p_head.add_run("iKHEDU")
            p_head.add_run("\t")
            r_hright = p_head.add_run(book_title)
            for run in (r_hleft, p_head.runs[1], r_hright):
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # B3.1: Watermark VML logo mờ 286pt giữa trang cho cả 3 header
        # (first / default / even) — Spec §2.1
        _add_vml_watermark(section.first_page_header, WATERMARK_LOGO)
        _add_vml_watermark(section.header, WATERMARK_LOGO)
        _add_vml_watermark(section.even_page_header, WATERMARK_LOGO)

        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="CBD5E1"/></w:pBdr>')
        p_head._p.get_or_add_pPr().append(pBdr)

        # Footer for page 2+ (Page numbering)
        footer = section.footer
        p_foot = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        run_f1 = p_foot.add_run("—  ")
        run_f1.font.name = "Times New Roman"
        run_f1.font.size = Pt(9)
        run_f1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
        p_foot._p.append(fldSimple)
        
        run_f2 = p_foot.add_run("  —")
        run_f2.font.name = "Times New Roman"
        run_f2.font.size = Pt(9)
        run_f2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ============================================================
    # 2. STYLES DEFINITION (Đồng bộ chuẩn 100% Level 1)
    # ============================================================
    style = doc.styles
    for style_name in ["Normal", "Body Text", "List Paragraph", "Compact", "First Paragraph"]:
        if style_name in [s.name for s in style]:
            st = style[style_name]
            st.font.name = "Times New Roman"
            st.font.size = Pt(12.5)
            st.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            st.paragraph_format.line_spacing = 1.1
            st.paragraph_format.space_after = Pt(2)

    # ============================================================
    # 3. (B2.1) KHÔNG CÒN TRANG BÌA — sách mở đầu thẳng vào Lời nói đầu.
    # ============================================================

    # ============================================================
    # 4. BOOKMARKS & PARAGRAPHS STYLING (Đồng bộ chuẩn 100% Level 1)
    # ============================================================
    toc_headings = []
    bm_counter = 1
    p_toc_heading = None
    in_preface = False      # B2.5: đang trong khối thân Lời nói đầu
    in_giaithich = False    # Khối dòng Giải thích của đề bài → căn trái, cấm Justify

    for idx, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else ""
        text = para.text.strip()

        # --- State machine cho Preface & Giải thích ---
        if style_name.startswith("Heading"):
            in_preface = False
            in_giaithich = False
        elif text.startswith("Giải thích:"):
            in_giaithich = True
        elif text.startswith(("Bối cảnh", "Đầu vào", "Đầu ra", "Ví dụ mẫu:", "Ràng buộc", "Bài ")):
            in_giaithich = False

        # --- A. HEADINGS ---
        if style_name == "Heading 1":
            is_chapter = text.startswith("CHƯƠNG")
            is_lesson = text.startswith("Bài")
            is_preface = "Lời nói đầu" in text
            is_appendix = text.startswith("Phụ lục")
            is_toc = "Mục lục" in text

            prev_text = doc.paragraphs[idx - 1].text.strip() if idx > 0 else ""
            is_first_lesson_in_chapter = is_lesson and prev_text.startswith("CHƯƠNG")

            # B2.5: P0 Lời nói đầu Center 14pt Bold, space_before 6pt
            if is_preface:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(6)
                para.paragraph_format.space_after = Pt(4)
                in_preface = True
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_chapter else WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.space_before = Pt(18) if is_chapter else (Pt(10) if is_first_lesson_in_chapter else Pt(16))
                para.paragraph_format.space_after = Pt(5) if is_chapter else Pt(4)
            para.paragraph_format.keep_with_next = True

            if not is_first_lesson_in_chapter and not is_preface:
                para._p.get_or_add_pPr().append(parse_xml(f'<w:pageBreakBefore {nsdecls("w")}/>'))

            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14) if is_preface else (Pt(18) if is_chapter else Pt(15.5))
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

            if is_toc:
                p_toc_heading = para
            else:
                bm_name = f"bm_sec_{bm_counter}"
                bm_counter += 1
                bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_counter}" w:name="{bm_name}"/>')
                bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_counter}"/>')
                para._p.insert(0, bm_start)
                para._p.append(bm_end)

                toc_headings.append({
                    "text": text,
                    "bm_name": bm_name,
                    "is_chapter": is_chapter,
                    "is_lesson": is_lesson,
                    "is_appendix": is_appendix,
                    "is_preface": is_preface,
                })

        elif style_name == "Heading 2":
            para.paragraph_format.space_before = Pt(11)
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.keep_with_next = True
            # B2.6: "Bài tập thực hành" → đỏ #FF0000 Bold 14pt
            h2_color = RGBColor(0xFF, 0x00, 0x00) if "Bài tập thực hành" in text else RGBColor(0x1E, 0x29, 0x3B)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = h2_color

        elif style_name == "Heading 3":
            para.paragraph_format.space_before = Pt(9)
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.keep_with_next = True
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        elif style_name == "Heading 4":
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.keep_with_next = True
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # --- B. TABLE CAPTION (Ví dụ/Bảng) - nổi bật, phân cấp rõ ràng ---
        # Chỉ nhận caption đánh số (Hình 4.1...), tránh nuốt dòng nội dung
        # bắt đầu bằng "Hình/Bảng" (vd: "Hình chữ nhật..." trong Giải thích)
        elif text.startswith("Ví dụ minh họa") or re.match(r"^(Bảng|Hình)\s+\d", text):
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(2)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # --- C. CODE BLOCKS ---
        elif "Code" in style_name or style_name == "Source Code":
            # B2.8: code Consolas 9pt / line 252, CĂN TRÁI tuyệt đối, xóa numPr rác
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for numPr in list(para._p.findall('.//' + qn('w:numPr'))):
                parent = numPr.getparent()
                if parent is not None:
                    parent.remove(numPr)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.05
            for r in para.runs:
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
            bdr = parse_xml(f'''<w:pBdr {nsdecls("w")}>
                <w:left w:val="single" w:sz="8" w:space="4" w:color="CBD5E1"/>
                <w:top w:val="single" w:sz="4" w:space="2" w:color="E2E8F0"/>
                <w:bottom w:val="single" w:sz="4" w:space="2" w:color="E2E8F0"/>
                <w:right w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
            </w:pBdr>''')
            ind = parse_xml(f'<w:ind {nsdecls("w")} w:left="180" w:right="120"/>')
            pPr = para._p.get_or_add_pPr()
            pPr.append(shd)
            pPr.append(bdr)
            pPr.append(ind)

        # --- C. CALLOUT / BLOCKQUOTES ---
        elif (style_name in ["Block Text", "Quote"] or text.startswith(">") or text.startswith("Định lý:") or text.startswith("Cảnh báo:")) and len(text) > 0:
            is_warning = "⚠️" in text or "Cảnh báo" in text or "Bẫy lỗi" in text
            is_tip = "💡" in text or "Mẹo" in text or "Lưu ý" in text
            is_theorem = "Định lý" in text or "Định nghĩa" in text or "Hệ quả" in text or "Bất biến" in text

            fill_color = "EFF6FF" if (is_warning or is_tip or is_theorem) else "F8FAFC"
            bdr_color = "3B82F6" if (is_warning or is_tip or is_theorem) else "94A3B8"

            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
            bdr = parse_xml(f'''<w:pBdr {nsdecls("w")}>
                <w:left w:val="single" w:sz="16" w:space="4" w:color="{bdr_color}"/>
                <w:top w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:pBdr>''')
            ind = parse_xml(f'<w:ind {nsdecls("w")} w:left="200" w:right="120"/>')
            pPr = para._p.get_or_add_pPr()
            pPr.append(shd)
            pPr.append(bdr)
            pPr.append(ind)

            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(3)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)

        # --- D. IMAGES & DRAWINGS ---
        elif len(para._p.findall('.//' + qn('w:drawing'))) > 0:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(2)
            
            for drawing in para._p.findall('.//' + qn('w:drawing')):
                wp_extent = drawing.find('.//' + qn('wp:extent'))
                a_ext = drawing.find('.//' + qn('a:ext'))
                if wp_extent is not None:
                    try:
                        old_cx = int(wp_extent.get('cx', 0))
                        old_cy = int(wp_extent.get('cy', 0))
                        if old_cx > 0 and old_cy > 0:
                            target_cx = 6000000
                            target_cy = int(round(old_cy * target_cx / old_cx))
                            if target_cy > 5000000:
                                target_cy = 5000000
                                target_cx = int(round(old_cx * target_cy / old_cy))
                            wp_extent.set('cx', str(target_cx))
                            wp_extent.set('cy', str(target_cy))
                            if a_ext is not None:
                                a_ext.set('cx', str(target_cx))
                                a_ext.set('cy', str(target_cy))
                    except Exception:
                        pass

        # --- E. IMAGE CAPTIONS ---
        elif style_name in ["Image Caption", "Caption"]:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(10)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.italic = True
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        # --- F. NORMAL TEXT (Fallback) ---
        else:
            if in_preface:
                # B2.5: thân Lời nói đầu Justify 14pt / 1.5
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.line_spacing = 1.5
                for r in para.runs:
                    if not r.font.name:
                        r.font.name = "Times New Roman"
                    r.font.size = Pt(14)
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            elif in_giaithich:
                # Dòng Giải thích: căn trái tự nhiên mọi style, cấm Justify (Spec §6.1)
                # (bao gồm cả List Paragraph kẻo dính Justify từ style gốc)
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.line_spacing = 1.1
                for r in para.runs:
                    if not r.font.name:
                        r.font.name = "Times New Roman"
                    r.font.size = Pt(12.5)
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            elif style_name in ("Body Text", "First Paragraph", "Compact", "Block Text"):
                # B2.7: Justify đúng 4 style này — CẤM Justify Normal
                # (Source Code kế thừa Normal, Justify sẽ làm dãn chữ code)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.line_spacing = 1.15
                for r in para.runs:
                    if not r.font.name:
                        r.font.name = "Times New Roman"
                    r.font.size = Pt(12.5)
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            else:
                if para.paragraph_format.space_after is None or para.paragraph_format.space_after.pt < 1:
                    para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.1
                for r in para.runs:
                    if not r.font.name:
                        r.font.name = "Times New Roman"
                    r.font.size = Pt(12.5)
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # ============================================================
    # 5. BUILD INTERACTIVE HYPERLINKED TOC AT END (Đồng bộ 100%)
    # ============================================================
    if p_toc_heading:
        print(f"  → Đang tạo Mục lục tương tác với {len(toc_headings)} liên kết...")
        last_elem = p_toc_heading._p
        
        for item in toc_headings:
            title = html.escape(item["text"])
            bm = item["bm_name"]
            
            if item["is_chapter"]:
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:spacing w:before="180" w:after="40"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:b/>
                                <w:sz w:val="23"/>
                                <w:color w:val="0F2A44"/>
                            </w:rPr>
                            <w:t>{title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            elif item["is_lesson"]:
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:ind w:left="280"/>
                        <w:spacing w:before="20" w:after="30"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:sz w:val="21"/>
                                <w:color w:val="1A4A6B"/>
                            </w:rPr>
                            <w:t>•  {title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            elif item["is_appendix"]:
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:spacing w:before="120" w:after="30"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:b/>
                                <w:sz w:val="22"/>
                                <w:color w:val="0F2A44"/>
                            </w:rPr>
                            <w:t>{title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            elif item["is_preface"]:
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:spacing w:before="100" w:after="40"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:b/>
                                <w:sz w:val="22"/>
                                <w:color w:val="0F2A44"/>
                            </w:rPr>
                            <w:t>{title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            else:
                continue
                
            p_elem = parse_xml(p_xml)
            last_elem.addnext(p_elem)
            last_elem = p_elem

    # ============================================================
    # 6. TABLES STYLING & BORDERS (Đồng bộ 100% Level 1)
    # ============================================================
    print(f"  → Đang định dạng và làm sạch viền cho {len(doc.tables)} bảng dữ liệu...")

    for s in doc.styles:
        if s.name == "Table":
            for el in s._element.findall(qn("w:tblStylePr")):
                s._element.remove(el)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        tblPr = table._tbl.tblPr
        for b in tblPr.findall(qn("w:tblBorders")):
            tblPr.remove(b)

        tblBorders = parse_xml(f'''<w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        </w:tblBorders>''')
        tblPr.append(tblBorders)

        # B3.2: xóa triệt để tblHeader (chống lặp tiêu đề khi ngắt trang)
        for th in table._tbl.findall('.//' + qn('w:tblHeader')):
            th_parent = th.getparent()
            if th_parent is not None:
                th_parent.remove(th)

        for r_idx, row in enumerate(table.rows):
            is_header = (r_idx == 0)
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            for c_idx, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                for b in tcPr.findall(qn("w:tcBorders")):
                    tcPr.remove(b)

                tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

                is_dense = len(row.cells) > 4
                header_font_size = Pt(9.5) if is_dense else Pt(10.5)
                math_sz = "19" if is_dense else "21"

                if is_header:
                    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.keep_with_next = True
                        for run in p.runs:
                            run.font.name = "Times New Roman"
                            run.font.bold = True
                            run.font.size = header_font_size
                            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

                        for r_elem in p._p.findall(".//" + qn("w:r")):
                            rPr = r_elem.get_or_add_rPr()
                            for c in rPr.findall(qn("w:color")):
                                rPr.remove(c)
                            rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="1E293B"/>'))

                        math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
                        for mr in p._p.findall(".//{" + math_ns + "}r"):
                            existing_rPr = mr.find(qn("w:rPr"))
                            if existing_rPr is not None:
                                for c in existing_rPr.findall(qn("w:color")):
                                    existing_rPr.remove(c)
                                existing_rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="1E293B"/>'))
                            else:
                                rPr = parse_xml(f'<w:rPr {nsdecls("w")}><w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:b {nsdecls("w")}/><w:color {nsdecls("w")} w:val="1E293B"/><w:sz {nsdecls("w")} w:val="{math_sz}"/></w:rPr>')
                                mr.insert(0, rPr)
                else:
                    fill_color = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
                    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>'))
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(1.5) if is_dense else Pt(2)
                        p.paragraph_format.space_after = Pt(1.5) if is_dense else Pt(2)
                        for run in p.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(9) if is_dense else Pt(10)
                            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

                tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="40" w:type="dxa"/><w:bottom w:w="40" w:type="dxa"/><w:left w:w="60" w:type="dxa"/><w:right w:w="60" w:type="dxa"/></w:tcMar>')
                tcPr.append(tcMar)

    # ============================================================
    # 7. BẢNG SAMPLE IO — B2.3 + Spec §5.3 (pass riêng sau style chung)
    # ============================================================
    print("  → Đang chuẩn hóa bảng Sample Input/Output (tblW 6800, Dynamic Indent)...")
    n_sample = 0
    for table in doc.tables:
        if len(table.rows) == 0 or len(table.rows[0].cells) != 2:
            continue
        if "Đầu vào (Input)" not in table.rows[0].cells[0].text:
            continue
        n_sample += 1

        tblPr = table._tbl.tblPr
        for tw in tblPr.findall(qn('w:tblW')):
            tblPr.remove(tw)
        tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="6800" w:type="dxa"/>'))

        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                for tcw in tcPr.findall(qn('w:tcW')):
                    tcPr.remove(tcw)
                tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="3400" w:type="dxa"/>'))

            if r_idx == 0:
                # Hàng tiêu đề: Consolas 11 Bold, căn giữa, nền F1F5F9
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            run.font.name = "Consolas"
                            run.font.size = Pt(11)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            else:
                # Hàng dữ liệu: tách <w:br/> → từng <w:p> độc lập (B2.3)
                for cell in row.cells:
                    for p in list(cell.paragraphs):
                        _split_para_at_br(p._p)
                # Dynamic Indent theo dòng dài nhất của hàng
                max_len = 0
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip():
                            max_len = max(max_len, len(p.text.rstrip()))
                indent = max(10.0, 72.0 - max(0, max_len - 4) * 3.25)
                for c_idx, cell in enumerate(row.cells):
                    for p_idx, p in enumerate(cell.paragraphs):
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_before = Pt(2) if p_idx == 0 else Pt(0)
                        p.paragraph_format.space_after = Pt(1)
                        p.paragraph_format.left_indent = Pt(indent)
                        for run in p.runs:
                            run.font.name = "Consolas"
                            run.font.size = Pt(11)
                            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    print(f"  → Đã chuẩn hóa {n_sample} bảng Sample.")

    # ============================================================
    # 8. B3.3 — MATH 12pt TOÀN TÀI LIỆU + BẢNG PHỤ LỤC A 12pt
    # ============================================================
    math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    n_math = 0
    for mr in doc.element.findall(f'.//{{{math_ns}}}r'):
        rPr = mr.find(qn('w:rPr'))
        if rPr is None:
            rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
            mr.insert(0, rPr)
        for tag in ('w:sz', 'w:szCs'):
            el = rPr.find(qn(tag))
            if el is None:
                el = parse_xml(f'<{tag} {nsdecls("w")} w:val="24"/>')
                rPr.append(el)
            else:
                el.set(qn('w:val'), '24')
        n_math += 1
    print(f"  → Đã ép {n_math} run công thức toán về 12pt.")

    # Bảng Phụ lục A → toàn bộ chữ 12pt (chạy sau pass Sample để override)
    in_app_a = False
    app_a_tables = set()
    for child in doc.element.body:
        if child.tag == qn('w:p'):
            t = (child.find(qn('w:pPr')) is not None)
            style_el = child.find(qn('w:pPr') + '/' + qn('w:pStyle')) if t else None
            is_h1 = style_el is not None and style_el.get(qn('w:val'), '').startswith('Heading1')
            if is_h1:
                txt = ''.join(n.text or '' for n in child.findall('.//' + qn('w:t')))
                in_app_a = txt.strip().startswith('Phụ lục A')
        elif child.tag == qn('w:tbl') and in_app_a:
            app_a_tables.add(child)
    n_app = 0
    for table in doc.tables:
        if table._tbl not in app_a_tables:
            continue
        n_app += 1
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)
    print(f"  → Đã đồng bộ {n_app} bảng Phụ lục A về 12pt.")

    doc.save(str(output_file))
    print(f"  ✅ Post-process Design System hoàn tất cho {output_file.name}!")


def _split_para_at_br(p_el):
    """B2.3: Tách một <w:p> thành nhiều <w:p> độc lập tại các run chỉ chứa <w:br/>."""
    from copy import deepcopy
    from docx.oxml.ns import qn as _qn
    groups = [[]]
    for child in list(p_el):
        if child.tag == _qn('w:pPr'):
            continue
        if child.tag == _qn('w:r') and child.find(_qn('w:br')) is not None and child.find(_qn('w:t')) is None:
            groups.append([])
        else:
            groups[-1].append(child)
    if len(groups) <= 1:
        return [p_el]
    for child in list(p_el):
        if child.tag != _qn('w:pPr'):
            p_el.remove(child)
    for child in groups[0]:
        p_el.append(child)
    result = [p_el]
    prev = p_el
    for g in groups[1:]:
        new_p = deepcopy(p_el)
        for child in list(new_p):
            if child.tag != _qn('w:pPr'):
                new_p.remove(child)
        for child in g:
            new_p.append(deepcopy(child))
        prev.addnext(new_p)
        prev = new_p
        result.append(new_p)
    return result


# ============================================================
# MAIN
# ============================================================

def build_single_volume(volume):
    if volume == 1:
        output_file = BASE_DIR / "IKHEDU_CPP_Nang_Cao_Quyen_1.docx"
    elif volume == 2:
        output_file = BASE_DIR / "IKHEDU_CPP_Nang_Cao_Quyen_2.docx"
    else:
        output_file = BASE_DIR / "IKHEDU_CPP_Nang_Cao_Xuat_Ban.docx"

    print("\n" + "=" * 60)
    print(f"📚 BẮT ĐẦU BUILD: {'QUYỂN ' + str(volume) if volume else 'BẢN TOÀN TẬP'}")
    print(f"   Đích xuất: {output_file.name}")
    print("=" * 60)

    markdown, book_title, book_subtitle = preprocess_markdown_for_volume(volume)
    success = build_with_pandoc(markdown, output_file)
    if not success:
        print(f"\n❌ Build thất bại cho {output_file.name}!")
        return False

    postprocess_docx(output_file, book_title, book_subtitle, volume)

    file_size = output_file.stat().st_size
    print(f"✅ HOÀN TẤT: {output_file.name} ({file_size / 1024 / 1024:.2f} MB)")
    return True

def main():
    parser = argparse.ArgumentParser(description="Build Word Docx for Level 2 Curriculum")
    parser.add_argument("--volume", type=int, choices=[1, 2], help="Build only Volume 1 or 2")
    parser.add_argument("--all", action="store_true", help="Build all Volumes (1,2)")
    args = parser.parse_args()

    if args.volume:
        build_single_volume(args.volume)
    else:
        print("=" * 60)
        print("📚 BẮT ĐẦU XUẤT BẢN BỘ SÁCH LEVEL 2 (346 BÀI TẬP — 2 QUYỂN CHUẨN IN ẤN)")
        print(f"   Tác giả: {BOOK_AUTHOR}")
        print("=" * 60)

        success1 = build_single_volume(1)
        success2 = build_single_volume(2)

        if success1 and success2:
            print("\n" + "=" * 60)
            print("🎉 TẤT CẢ CÁC QUYỂN LEVEL 2 ĐÃ ĐƯỢC XUẤT BẢN THÀNH CÔNG!")
            print("   📘 Quyển 1: IKHEDU_CPP_Nang_Cao_Quyen_1.docx (Chương 01-03, Bài 01-06 — 134 bài tập)")
            print("   📕 Quyển 2: IKHEDU_CPP_Nang_Cao_Quyen_2.docx (Chương 04-06, Bài 07-15 — 212 bài tập)")
            print("   ✨ Tổng quy mô: 346 bài toán thực hành chuẩn thi đấu (100% C++ Solution)!")
            print("=" * 60)

if __name__ == "__main__":
    main()
