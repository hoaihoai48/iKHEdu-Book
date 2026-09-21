import sys
import zipfile
import re
from pathlib import Path
from docx import Document

def audit_scratch_docx(docx_path):
    print(f"\n========================================================")
    print(f"🔍 AUDIT KIỂM ĐỊNH FILE: {docx_path.name}")
    print(f"========================================================")
    
    doc = Document(str(docx_path))
    
    # 1. Check Margins & Sections
    sec = doc.sections[0]
    print(f"1. Khổ giấy & Margins:")
    print(f"   - Top: {sec.top_margin.pt} pt (Chuẩn: 36.0 pt)")
    print(f"   - Bottom: {sec.bottom_margin.pt} pt (Chuẩn: 36.0 pt)")
    print(f"   - Left (Gáy): {sec.left_margin.pt} pt (Chuẩn: 64.35 pt)")
    print(f"   - Right: {sec.right_margin.pt} pt (Chuẩn: 36.0 pt)")
    print(f"   - TitlePg (different_first_page): {sec.different_first_page_header_footer}")
    
    # 2. Check Watermark in headers via zip
    with zipfile.ZipFile(str(docx_path), "r") as z:
        header_files = [f for f in z.namelist() if f.startswith("word/header")]
        print(f"2. Watermark Logo:")
        print(f"   - Các file header: {header_files}")
        watermark_count = 0
        for hf in header_files:
            content = z.read(hf).decode("utf-8", errors="ignore")
            if "WordPictureWatermark" in content or "logo_in" in content:
                watermark_count += 1
                print(f"     + {hf}: ĐÃ CÓ VML Watermark shape ✅")
            else:
                print(f"     + {hf}: CHƯA CÓ Watermark ❌")
        
        # Check settings.xml for updateFields
        if "word/settings.xml" in z.namelist():
            settings = z.read("word/settings.xml").decode("utf-8", errors="ignore")
            if "updateFields" in settings:
                print(f"   - updateFields: PHÁT HIỆN THẺ NGUY HIỂM ❌")
            else:
                print(f"   - updateFields: Sạch 100% (An toàn bảo mật) ✅")

    # 3. Check Tables
    total_tables = len(doc.tables)
    bad_headers = 0
    sample_tables = 0
    cantsplit_rows = 0
    total_rows = 0
    for t in doc.tables:
        for r in t.rows:
            total_rows += 1
            tr_xml = r._tr.xml
            if "w:tblHeader" in tr_xml:
                bad_headers += 1
            if "w:cantSplit" in tr_xml:
                cantsplit_rows += 1
        if len(t.rows) >= 2 and len(t.rows[0].cells) == 2:
            t_txt = t.rows[0].cells[0].text + t.rows[0].cells[1].text
            if "Input" in t_txt or "Đầu vào" in t_txt:
                sample_tables += 1

    print(f"3. Bảng dữ liệu (Tables):")
    print(f"   - Tổng số bảng: {total_tables}")
    print(f"   - Bảng Sample IO: {sample_tables}")
    print(f"   - Hàng mang tblHeader (cấm > 0): {bad_headers} (Chuẩn: 0) {'✅' if bad_headers == 0 else '❌'}")
    print(f"   - Hàng mang cantSplit: {cantsplit_rows}/{total_rows} {'✅' if cantsplit_rows == total_rows else '⚠️'}")

    # 4. Check Colors & Heading 2 Bài tập thực hành
    red_count = 0
    practice_h2 = 0
    sample_io_colors = []
    for p in doc.paragraphs:
        if p.style and p.style.name == "Heading 2" and "Bài tập thực hành" in p.text:
            practice_h2 += 1
            for r in p.runs:
                if r.font.color and r.font.color.rgb:
                    col = str(r.font.color.rgb)
                    if col.upper() == "FF0000":
                        red_count += 1

    print(f"4. Màu sắc & Điểm nhấn thị giác:")
    print(f"   - Heading 2 'Bài tập thực hành': {practice_h2} mục")
    print(f"   - Số mục mang màu đỏ #FF0000 chuẩn: {red_count}/{practice_h2} {'✅' if red_count == practice_h2 else '❌'}")

    # 5. Check TOC
    toc_paragraphs = [p for p in doc.paragraphs if "MỤC LỤC" in p.text or "Mục lục" in p.text]
    print(f"5. Mục lục tương tác (TOC):")
    print(f"   - Tìm thấy tiêu đề Mục lục: {len(toc_paragraphs)} mục {'✅' if len(toc_paragraphs) > 0 else '❌'}")
    if toc_paragraphs:
        tp = toc_paragraphs[0]
        for r in tp.runs:
            print(f"   - Cỡ chữ tiêu đề Mục lục: {r.font.size.pt} pt (Chuẩn: 18.0 pt)")

if __name__ == "__main__":
    base = Path("courses/scratch-bang-a")
    for f in ["scratch-quyen-1.docx", "scratch-quyen-2.docx"]:
        p = base / f
        if p.exists():
            audit_scratch_docx(p)
