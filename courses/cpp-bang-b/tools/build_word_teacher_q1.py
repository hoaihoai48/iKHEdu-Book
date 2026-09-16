#!/usr/bin/env python3
"""
BUILD SCRIPT: SÁCH GIÁO VIÊN C++ QUYỂN 1 (CHƯƠNG 01 - 04, BÀI 01 - 12, 188 BÀI TOÁN)
Chuẩn khuôn 100% theo python-giaovien-full.docx:
1. Hình ảnh minh họa trực quan bối cảnh (Pure Scenario Illustration) gắn ngay sau Bối cảnh của từng bài.
2. Tiêu đề ảnh italic (Times New Roman 10pt, màu #475569).
3. Bảng Sample IO: Header nền #F1F5F9, chữ Consolas 11pt Bold màu #000000, data cell nền trắng #FFFFFF với Dynamic Left Indent (72pt -> 10pt).
4. Đệm cách trước và sau mọi bảng: paragraph trước w:after="120" (6pt), paragraph sau w:before="120" (6pt).
5. Khung Code: Consolas 9pt, line 1.05, nền #F8FAFC, viền #E2E8F0, căn trái 100%.
6. Toàn bộ chữ Body và Headings dùng màu đen tuyền #000000 (độ tương phản sắc nét in ấn).
7. Header trang: Tab Stop right w:pos="9899", Watermark logo image13.jpeg đầy đủ 3 headers (default, even, first).
8. Mục lục: Hyperlink tinh gọn 4 Chương + 12 Bài học vừa vặn 1 trang, không vỡ trang.
"""

import os
import re
import sys
import json
import html
import shutil
import zipfile
import subprocess
from pathlib import Path

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path("/Users/vu/Developer/ikhEdu_lessons/courses/cpp-bang-b")
PROBLEMS_DIR = BASE_DIR / "problems"
MANIFEST_FILE = BASE_DIR / "word_build_manifest_gv.json"
PREFACE_FILE = BASE_DIR / "reference" / "Loi_Noi_Dau_GV.md"
TEMPLATE_DOCX = BASE_DIR / "c++-level-1-quyen-1.docx"
OUT_DOCX = BASE_DIR / "cpp-giaovien-quyen-1.docx"
DIAGRAMS_MAP_FILE = BASE_DIR / "problem_diagrams_q1_map.json"
ASSETS_PNG_DIR = BASE_DIR / "assets_png"

CHAPTER_NAMES = {
    1: "CHƯƠNG 01: THUẬT TOÁN SẮP XẾP & KỸ THUẬT MẢNG",
    2: "CHƯƠNG 02: MẢNG TIỀN TỐ & TÌM KIẾM NHỊ PHÂN",
    3: "CHƯƠNG 03: SỐ HỌC & ĐẠI SỐ MODULAR",
    4: "CHƯƠNG 04: ĐỆ QUY, CHIA ĐỂ TRỊ & QUAY LUI",
}

# Load problem illustrations map
if DIAGRAMS_MAP_FILE.exists():
    with open(DIAGRAMS_MAP_FILE, "r", encoding="utf-8") as f:
        DIAGRAMS_MAP = json.load(f)
else:
    DIAGRAMS_MAP = {}

def parse_de_bai(de_bai_path):
    if not de_bai_path.exists():
        return {"bcanh": "", "nv": "", "inp": "", "out": "", "samples": []}
    text = de_bai_path.read_text(encoding="utf-8")

    bcanh_m = re.search(r'## Bối cảnh\s*\n(.*?)(?=\n## Nhiệm vụ|\n## Input|\Z)', text, re.DOTALL)
    nv_m = re.search(r'## Nhiệm vụ\s*\n(.*?)(?=\n## Input|\n## Output|\Z)', text, re.DOTALL)
    inp_m = re.search(r'## Input\s*\n(.*?)(?=\n## Output|\n## Sample|\Z)', text, re.DOTALL)
    out_m = re.search(r'## Output\s*\n(.*?)(?=\n## Sample|\n## Ràng buộc|\Z)', text, re.DOTALL)

    bcanh = bcanh_m.group(1).strip() if bcanh_m else ""
    nv = nv_m.group(1).strip() if nv_m else ""
    inp = inp_m.group(1).strip() if inp_m else ""
    out = out_m.group(1).strip() if out_m else ""

    samples = []
    sample_blocks = re.findall(r'## (Sample \d+)\s*\n(.*?)(?=\n## Sample \d+|\n## Ràng buộc|\Z)', text, re.DOTALL)
    for s_title, s_body in sample_blocks:
        sinp_m = re.search(r'### Input\s*```(?:text)?\n(.*?)```', s_body, re.DOTALL)
        sout_m = re.search(r'### Output\s*```(?:text)?\n(.*?)```', s_body, re.DOTALL)
        sexpl_m = re.search(r'### Giải thích\s*\n(.*)', s_body, re.DOTALL)

        sinp = sinp_m.group(1).strip() if sinp_m else ""
        sout = sout_m.group(1).strip() if sout_m else ""
        sexpl = sexpl_m.group(1).strip() if sexpl_m else ""

        samples.append({
            "title": s_title,
            "input": sinp,
            "output": sout,
            "explain": sexpl
        })

    return {
        "bcanh": bcanh,
        "nv": nv,
        "inp": inp,
        "out": out,
        "samples": samples
    }

def parse_guide(guide_path):
    if not guide_path.exists():
        return []
    text = guide_path.read_text(encoding="utf-8")

    sections = []
    raw_secs = re.split(r'\n(?=##\s+\d+\.)', text)

    for raw in raw_secs:
        m = re.match(r'##\s+(\d+\.\s+[^\n]+)\n(.*)', raw.strip(), re.DOTALL)
        if m:
            sec_title = m.group(1).strip()
            sec_body = m.group(2).strip()
            sec_body = re.sub(r'\n---\s*$', '', sec_body).strip()
            sec_body = re.sub(r'^(#{1,3})\s+', '##### ', sec_body, flags=re.MULTILINE)
            sections.append({
                "title": sec_title,
                "body": sec_body
            })

    return sections

def format_teacher_problem_markdown(prob_idx, prob_code, prob_title, parsed_de_bai, sol_code, parsed_guide):
    md = []
    md.append(f"### Bài {prob_idx:02d} [{prob_code}]: {prob_title}\n\n")

    if parsed_de_bai["bcanh"]:
        md.append(f"Bối cảnh: {parsed_de_bai['bcanh']}\n\n")

    if parsed_de_bai["nv"]:
        md.append(f"Nhiệm vụ: {parsed_de_bai['nv']}\n\n")

    if parsed_de_bai["inp"]:
        md.append(f"**Đầu vào (Input):**\n\n{parsed_de_bai['inp']}\n\n")

    if parsed_de_bai["out"]:
        md.append(f"**Đầu ra (Output):**\n\n{parsed_de_bai['out']}\n\n")

    for s in parsed_de_bai["samples"]:
        s_title_str = f"Ví dụ mẫu ({s['title']}):"
        inp_str = s["input"].replace("\n", " <br> ")
        out_str = s["output"].replace("\n", " <br> ")

        md.append(f"**{s_title_str}**\n\n")
        md.append("| Đầu vào (Input) | Đầu ra (Output) |\n")
        md.append("| :--- | :--- |\n")
        md.append(f"| {inp_str} | {out_str} |\n\n")

        if s["explain"]:
            md.append(f"**Giải thích:** {s['explain']}\n\n")

    for g_sec in parsed_guide:
        if "Lời giải tham khảo" in g_sec["title"]:
            continue
        md.append(f"#### {g_sec['title']}\n\n")
        md.append(f"{g_sec['body']}\n\n")

    md.append("#### 4. Lời giải tham khảo\n\n")
    md.append(f"```cpp\n{sol_code}\n```\n\n")

    return "".join(md)

def build_q1_markdown(manifest_data):
    md_parts = []
    
    # Lời nói đầu (Heading 1)
    md_parts.append("# Lời nói đầu\n\n")
    md_parts.append(
        "Cuốn **Giáo trình C++ — Sách giáo viên — Quyển 1** được biên soạn đồng bộ cùng chương trình đào tạo "
        "lập trình thuật toán C++ cơ bản iKHEDU, dành riêng cho quý thầy cô giáo và các huấn luyện viên chuyên trách. "
        "Tài liệu này cung cấp toàn bộ lời giải chi tiết, chuẩn mực và hệ thống **hướng dẫn sư phạm thực chiến chuyên sâu** "
        "(ý tưởng và phân tích thuật toán, bảng mô phỏng chạy tay từng bước trên số liệu mẫu thực tế, lưu ý và các bẫy lỗi kinh điển thường gặp) "
        "cho toàn bộ **188 bài toán thực hành** thuộc 4 Chương đầu (Bài 01 đến Bài 12).\n\n"
        "Mỗi bài toán trong sách đều được cấu trúc thống nhất từ đề bài học sinh, phân tích tư duy giải pháp, bảng theo dõi biến số trực quan "
        "đến mã nguồn tham khảo tối ưu nhất (Fast I/O, Safe Input, 0 std::), giúp thầy cô dễ dàng tổ chức giờ học chất lượng cao và đồng hành hiệu quả cùng học sinh.\n\n"
    )

    current_chapter = None
    q1_lessons = [l for l in manifest_data["lessons"] if l["chapter"] in [1, 2, 3, 4]]
    prob_count = 0

    for lesson_info in q1_lessons:
        chap_num = lesson_info["chapter"]
        chap_title = CHAPTER_NAMES.get(chap_num, f"CHƯƠNG {chap_num:02d}")
        lesson_title = lesson_info["title"]

        if chap_num != current_chapter:
            current_chapter = chap_num
            md_parts.append(f"# {chap_title}\n\n")

        md_parts.append(f"# {lesson_title}\n\n")

        for p_idx, prob in enumerate(lesson_info["problems"], 1):
            prob_code = prob["code"]
            prob_title = prob["title"]
            de_bai_file = BASE_DIR / prob["de_bai"]
            sol_file = BASE_DIR / prob["solution"]
            guide_file = BASE_DIR / prob["guide"]

            parsed_de_bai = parse_de_bai(de_bai_file)
            sol_code = sol_file.read_text(encoding="utf-8").strip() if sol_file.exists() else ""
            parsed_guide = parse_guide(guide_file)

            p_md = format_teacher_problem_markdown(p_idx, prob_code, prob_title, parsed_de_bai, sol_code, parsed_guide)
            md_parts.append(p_md)
            prob_count += 1

    md_parts.append("# Mục lục\n\n")
    print(f"  → Tổng hợp Markdown Quyển 1: {len(q1_lessons)} bài học, {prob_count} bài toán.")
    return "".join(md_parts)

def convert_markdown_to_docx(full_md, out_docx):
    print("🔧 Đang chạy Pandoc Markdown → DOCX...")
    intermediate_md = BASE_DIR / "_build_intermediate_gv_q1.md"
    intermediate_md.write_text(full_md, encoding="utf-8")

    cmd = [
        "pandoc",
        str(intermediate_md),
        "-o", str(out_docx),
        "--reference-doc", str(TEMPLATE_DOCX),
        "--from=markdown+tex_math_dollars+pipe_tables+fenced_code_blocks+yaml_metadata_block+header_attributes",
        "--to=docx",
        "--standalone",
        "--wrap=auto",
        "--columns=80"
    ]

    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"❌ Pandoc lỗi: {res.stderr}")
        sys.exit(1)

    print(f"  ✅ Pandoc thành công: {out_docx.name} ({out_docx.stat().st_size:,} bytes)")
    if intermediate_md.exists():
        intermediate_md.unlink()
    return True

def insert_picture_with_caption(doc, img_filename, caption_text, target_p_elem):
    """Chèn hình minh họa bối cảnh centered và chú thích italic vào trước target_p_elem."""
    img_path = ASSETS_PNG_DIR / img_filename
    if not img_path.exists():
        return

    # 1. Paragraph chứa hình ảnh
    p_pic = doc.add_paragraph()
    p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pic.paragraph_format.space_before = Pt(6)
    p_pic.paragraph_format.space_after = Pt(2)
    p_pic.paragraph_format.line_spacing = 1.0
    r_pic = p_pic.add_run()
    r_pic.add_picture(str(img_path), width=Inches(5.0))
    target_p_elem.addprevious(p_pic._p)

    # 2. Paragraph chứa caption
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(6)
    p_cap.paragraph_format.line_spacing = 1.0
    r_cap = p_cap.add_run(f"Hình minh họa: {caption_text}")
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    target_p_elem.addprevious(p_cap._p)

def post_process_teacher_docx(docx_path, manifest_data):
    print("🎨 Đang áp dụng thiết kế Design System chuẩn khuôn Python...")
    doc = docx.Document(str(docx_path))
    s = doc.sections[0]

    # 1. Margins
    s.top_margin = Pt(36.0)
    s.bottom_margin = Pt(36.0)
    s.left_margin = Pt(64.35)  # gáy sách
    s.right_margin = Pt(36.0)

    # 2. Update Headers (3 headers: even, default, first)
    header_book_text = "Giáo trình C++ - Sách giáo viên"
    for h in [s.even_page_header, s.header, s.first_page_header]:
        p0 = h.paragraphs[0]
        pPr = p0._p.get_or_add_pPr()
        for tb in pPr.findall(qn("w:tabs")): pPr.remove(tb)
        pPr.append(parse_xml(f'<w:tabs {nsdecls("w")}><w:tab {nsdecls("w")} w:val="right" w:pos="9899"/></w:tabs>'))

        watermark_runs = []
        for r in p0.runs:
            if r._r.find(qn("w:pict")) is not None:
                watermark_runs.append(r._r)

        for child in list(p0._p):
            if child.tag.split("}")[-1] != "pPr":
                p0._p.remove(child)

        for wr in watermark_runs:
            p0._p.append(wr)

        r_left = parse_xml(f'''<w:r {nsdecls("w")}>
            <w:rPr>
                <w:rFonts w:cs="Times New Roman"/>
                <w:b/><w:bCs/><w:i/><w:iCs/>
                <w:sz w:val="22"/><w:szCs w:val="22"/>
            </w:rPr>
            <w:t>Trung tâm giáo dục &amp; đào tạo tin học iKH</w:t>
        </w:r>''')
        p0._p.append(r_left)

        r_tab = parse_xml(f'''<w:r {nsdecls("w")}>
            <w:rPr><w:rFonts w:cs="Times New Roman"/><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr>
            <w:tab/>
        </w:r>''')
        p0._p.append(r_tab)

        r_right = parse_xml(f'''<w:r {nsdecls("w")}>
            <w:rPr>
                <w:rFonts w:cs="Times New Roman"/>
                <w:i/><w:iCs/><w:noProof/>
                <w:sz w:val="22"/><w:szCs w:val="22"/>
                <w:lang w:val="vi-VN"/>
            </w:rPr>
            <w:t>{html.escape(header_book_text)}</w:t>
        </w:r>''')
        p0._p.append(r_right)

    # 3. Process Paragraphs, Headings & Bookmarks
    toc_headings = []
    p_toc_heading = None

    for idx, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else ""
        text = para.text.strip()

        # Heading 1
        if style_name == "Heading 1" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading1"):
            para.paragraph_format.keep_with_next = True
            is_chapter = "CHƯƠNG" in text.upper()
            is_preface = "LỜI NÓI ĐẦU" in text.upper()
            is_toc = "MỤC LỤC" in text.upper()

            if is_toc:
                p_toc_heading = para
                para.paragraph_format.page_break_before = True
                continue

            if is_chapter:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(18)
                para.paragraph_format.space_after = Pt(5)
                para.paragraph_format.page_break_before = True
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(18)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            elif is_preface:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(6)
                para.paragraph_format.space_after = Pt(6)
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(14)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            else:
                # Lesson H1
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.space_before = Pt(14)
                para.paragraph_format.space_after = Pt(4)
                prev_text = doc.paragraphs[idx - 1].text.strip() if idx > 0 else ""
                if prev_text.startswith("CHƯƠNG"):
                    para.paragraph_format.page_break_before = False
                else:
                    para.paragraph_format.page_break_before = True

                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(15.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

            if not is_preface and text:
                bm_name = f"_toc_h1_{idx}"
                para._p.insert(0, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{idx}" w:name="{bm_name}"/>'))
                para._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{idx}"/>'))
                toc_headings.append({
                    "text": text,
                    "bm_name": bm_name,
                    "level": 1,
                    "is_chapter": is_chapter,
                    "is_lesson": not is_chapter
                })

        # Heading 2
        elif style_name == "Heading 2" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading2"):
            para.paragraph_format.keep_with_next = True
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(4)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Heading 3 (Tên bài toán)
        elif style_name == "Heading 3" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading3"):
            para.paragraph_format.keep_with_next = True
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(10)
            para.paragraph_format.space_after = Pt(4)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Heading 4 (4 Mục Hướng dẫn)
        elif style_name == "Heading 4" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading4"):
            para.paragraph_format.keep_with_next = True
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(3)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Heading 5
        elif style_name == "Heading 5" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading5"):
            para.paragraph_format.keep_with_next = True
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(2)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Source Code
        elif style_name == "Source Code" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "SourceCode"):
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pPr = para._p.get_or_add_pPr()
            for numPr in pPr.findall(qn("w:numPr")): pPr.remove(numPr)
            for r in para.runs:
                r.font.name = "Consolas"
                r.font.size = Pt(9.0)
                r.font.color.rgb = RGBColor(0x0F, 0x2A, 0x44)

            sp = pPr.find(qn("w:spacing"))
            if sp is not None:
                sp.set(qn("w:before"), "20")
                sp.set(qn("w:after"), "20")
                sp.set(qn("w:line"), "252")
                sp.set(qn("w:lineRule"), "auto")
            else:
                pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="20" w:after="20" w:line="252" w:lineRule="auto"/>'))

            if pPr.find(qn("w:shd")) is None:
                pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
            if pPr.find(qn("w:pBdr")) is None:
                bdr = parse_xml(f'''<w:pBdr {nsdecls("w")}>
                    <w:left w:val="single" w:sz="6" w:space="4" w:color="CBD5E1"/>
                    <w:top w:val="single" w:sz="4" w:space="2" w:color="E2E8F0"/>
                    <w:bottom w:val="single" w:sz="4" w:space="2" w:color="E2E8F0"/>
                    <w:right w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                </w:pBdr>''')
                pPr.append(bdr)
            if pPr.find(qn("w:ind")) is None:
                pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="180" w:right="120"/>'))

        # Callout / Quote
        elif style_name in ["Block Text", "Quote"] or text.startswith(">"):
            pPr = para._p.get_or_add_pPr()
            if pPr.find(qn("w:shd")) is None:
                pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="EFF6FF"/>'))
            if pPr.find(qn("w:pBdr")) is None:
                bdr = parse_xml(f'''<w:pBdr {nsdecls("w")}>
                    <w:left w:val="single" w:sz="16" w:space="4" w:color="3B82F6"/>
                    <w:top w:val="single" w:sz="4" w:space="3" w:color="E2E8F0"/>
                    <w:bottom w:val="single" w:sz="4" w:space="3" w:color="E2E8F0"/>
                    <w:right w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                </w:pBdr>''')
                pPr.append(bdr)
            if pPr.find(qn("w:ind")) is None:
                pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="200" w:right="120"/>'))

            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(3)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Normal / Compact / Body Text
        else:
            if idx <= 2:
                # Preface body
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.line_spacing = 1.5
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(14)
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            else:
                if text.startswith("Giải thích:") or text.startswith("**Giải thích:**"):
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.line_spacing = 1.15
                for r in para.runs:
                    if not r.font.name:
                        r.font.name = "Times New Roman"
                    r.font.size = Pt(12.5)
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 4. Gắn hình ảnh minh họa bối cảnh vào từng bài toán
    print(f"🖼️ Đang chèn {len(DIAGRAMS_MAP)} hình ảnh minh họa trực quan bối cảnh...")
    for prob_code, (img_filename, caption_text) in DIAGRAMS_MAP.items():
        h3_idx = None
        for i, p in enumerate(doc.paragraphs):
            if prob_code in p.text and p.style.name == "Heading 3":
                h3_idx = i
                break
        if h3_idx is None:
            continue

        target_p = None
        for j in range(h3_idx + 1, min(len(doc.paragraphs), h3_idx + 6)):
            p_curr = doc.paragraphs[j]
            if p_curr.style.name == "Heading 3":
                break
            if p_curr.text.strip().startswith("Bối cảnh:"):
                next_p = doc.paragraphs[j + 1] if j + 1 < len(doc.paragraphs) else p_curr
                target_p = next_p
                break
        if target_p is None:
            target_p = doc.paragraphs[h3_idx + 1] if h3_idx + 1 < len(doc.paragraphs) else doc.paragraphs[h3_idx]

        insert_picture_with_caption(doc, img_filename, caption_text, target_p._p)

    # 5. Tables Processing & Dynamic Left Indent & Spacing 120 dxa
    body_element = doc._body._element
    for i, child in enumerate(body_element):
        if child.tag.split("}")[-1] == "tbl":
            prev_el = body_element[i - 1] if i > 0 else None
            if prev_el is not None and prev_el.tag.split("}")[-1] == "p":
                prev_pPr = prev_el.find(qn("w:pPr"))
                if prev_pPr is not None:
                    sp = prev_pPr.find(qn("w:spacing"))
                    if sp is not None:
                        curr_after = sp.attrib.get(qn("w:after"), None)
                        if curr_after is None or int(curr_after) < 120:
                            sp.set(qn("w:after"), "120")
                    else:
                        prev_pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:after="120"/>'))

            next_el = body_element[i + 1] if i + 1 < len(body_element) else None
            if next_el is not None and next_el.tag.split("}")[-1] == "p":
                next_pPr = next_el.find(qn("w:pPr"))
                if next_pPr is not None:
                    sp = next_pPr.find(qn("w:spacing"))
                    if sp is not None:
                        curr_before = sp.attrib.get(qn("w:before"), None)
                        if curr_before is None or int(curr_before) < 120:
                            sp.set(qn("w:before"), "120")
                    else:
                        next_pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="120"/>'))

    for table in doc.tables:
        table_tblPr = table._tbl.tblPr

        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            for h in trPr.findall(qn("w:tblHeader")): trPr.remove(h)
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        for mr in table._tbl.findall(".//{" + math_ns + "}r"):
            mr_rPr = mr.find(qn("w:rPr"))
            if mr_rPr is not None:
                for sz in mr_rPr.findall(qn("w:sz")): sz.set(qn("w:val"), "24")
                for szCs in mr_rPr.findall(qn("w:szCs")): szCs.set(qn("w:val"), "24")
            else:
                mr.insert(0, parse_xml(f'<w:rPr {nsdecls("w")}><w:sz {nsdecls("w")} w:val="24"/><w:szCs {nsdecls("w")} w:val="24"/></w:rPr>'))

        is_sample = False
        if len(table.rows) >= 2 and len(table.rows[0].cells) == 2:
            c0_t = table.rows[0].cells[0].text.strip()
            c1_t = table.rows[0].cells[1].text.strip()
            if ("Đầu vào" in c0_t or "Input" in c0_t) and ("Đầu ra" in c1_t or "Output" in c1_t):
                is_sample = True

        if is_sample:
            table_tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
            tblW = table_tblPr.find(qn("w:tblW"))
            if tblW is not None:
                tblW.set(qn("w:w"), "6800")
                tblW.set(qn("w:type"), "dxa")
            else:
                table_tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="6800" w:type="dxa"/>'))

            for b in table_tblPr.findall(qn("w:tblBorders")): table_tblPr.remove(b)
            table_tblPr.append(parse_xml(f'''<w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            </w:tblBorders>'''))

            # Row 0: Headers
            for cell in table.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                for w in tcPr.findall(qn("w:tcW")): tcPr.remove(w)
                tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="3400" w:type="dxa"/>'))
                for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
                for tm in tcPr.findall(qn("w:tcMar")): tcPr.remove(tm)
                tcPr.append(parse_xml(f'''<w:tcMar {nsdecls("w")}>
                    <w:top w:w="40" w:type="dxa"/>
                    <w:left w:w="60" w:type="dxa"/>
                    <w:bottom w:w="40" w:type="dxa"/>
                    <w:right w:w="60" w:type="dxa"/>
                </w:tcMar>'''))
                for va in tcPr.findall(qn("w:vAlign")): tcPr.remove(va)
                tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

                for p in cell.paragraphs:
                    pPr = p._p.get_or_add_pPr()
                    sp = pPr.find(qn("w:spacing"))
                    if sp is not None:
                        sp.set(qn("w:before"), "60")
                        sp.set(qn("w:after"), "60")
                        sp.set(qn("w:line"), "276")
                        sp.set(qn("w:lineRule"), "auto")
                    else:
                        pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60" w:line="276" w:lineRule="auto"/>'))
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Consolas"
                        r.font.size = Pt(11)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

            # Row 1+: Testcase data
            for r_idx in range(1, len(table.rows)):
                row = table.rows[r_idx]
                for cell in row.cells:
                    tcPr = cell._tc.get_or_add_tcPr()
                    for w in tcPr.findall(qn("w:tcW")): tcPr.remove(w)
                    tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="3400" w:type="dxa"/>'))
                    for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>'))
                    for tm in tcPr.findall(qn("w:tcMar")): tcPr.remove(tm)
                    tcPr.append(parse_xml(f'''<w:tcMar {nsdecls("w")}>
                        <w:top w:w="40" w:type="dxa"/>
                        <w:left w:w="60" w:type="dxa"/>
                        <w:bottom w:w="40" w:type="dxa"/>
                        <w:right w:w="60" w:type="dxa"/>
                    </w:tcMar>'''))
                    for va in tcPr.findall(qn("w:vAlign")): tcPr.remove(va)
                    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

                    orig_p_list = list(cell.paragraphs)
                    lines = []
                    for p in orig_p_list:
                        p_text = p.text
                        if "<br>" in p_text or "<br/>" in p_text or "&lt;br&gt;" in p_text:
                            sub = re.split(r'<br\s*/?>|&lt;br&gt;', p_text)
                            for s_str in sub:
                                s_clean = s_str.strip()
                                if s_clean: lines.append(s_clean)
                        else:
                            p_clean = p_text.strip()
                            if p_clean: lines.append(p_clean)

                    if not lines: lines = [""]

                    max_len = max(len(l) for l in lines) if lines else 0
                    if max_len <= 4:
                        dyn_indent_pt = 72.0
                    else:
                        dyn_indent_pt = max(10.0, 72.0 - (max_len - 4) * 3.25)
                    dyn_indent_dxa = int(round(dyn_indent_pt * 20))

                    for p in orig_p_list:
                        p_elem = p._p
                        if p_elem.getparent() is not None:
                            p_elem.getparent().remove(p_elem)

                    for line_str in lines:
                        new_p = parse_xml(f'''<w:p {nsdecls("w")}>
                            <w:pPr>
                                <w:jc w:val="left"/>
                                <w:ind w:left="{dyn_indent_dxa}"/>
                                <w:spacing w:before="30" w:after="30" w:line="240" w:lineRule="auto"/>
                            </w:pPr>
                            <w:r>
                                <w:rPr>
                                    <w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>
                                    <w:sz w:val="19"/>
                                    <w:szCs w:val="19"/>
                                    <w:color w:val="0F2A44"/>
                                </w:rPr>
                                <w:t xml:space="preserve">{html.escape(line_str)}</w:t>
                            </w:r>
                        </w:p>''')
                        cell._tc.append(new_p)

        else:
            # Dry Run Table
            table_tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
            for b in table_tblPr.findall(qn("w:tblBorders")): table_tblPr.remove(b)
            table_tblPr.append(parse_xml(f'''<w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            </w:tblBorders>'''))

            for r_idx, row in enumerate(table.rows):
                for cell in row.cells:
                    tcPr = cell._tc.get_or_add_tcPr()
                    for tm in tcPr.findall(qn("w:tcMar")): tcPr.remove(tm)
                    tcPr.append(parse_xml(f'''<w:tcMar {nsdecls("w")}>
                        <w:top w:w="60" w:type="dxa"/>
                        <w:left w:w="80" w:type="dxa"/>
                        <w:bottom w:w="60" w:type="dxa"/>
                        <w:right w:w="80" w:type="dxa"/>
                    </w:tcMar>'''))
                    for va in tcPr.findall(qn("w:vAlign")): tcPr.remove(va)
                    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

                    if r_idx == 0:
                        for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
                        for p in cell.paragraphs:
                            pPr = p._p.get_or_add_pPr()
                            sp = pPr.find(qn("w:spacing"))
                            if sp is not None:
                                sp.set(qn("w:before"), "60")
                                sp.set(qn("w:after"), "60")
                                sp.set(qn("w:line"), "276")
                                sp.set(qn("w:lineRule"), "auto")
                            else:
                                pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60" w:line="276" w:lineRule="auto"/>'))
                            for r in p.runs:
                                r.font.name = "Times New Roman"
                                r.font.size = Pt(12)
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
                    else:
                        fill_c = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
                        for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_c}"/>'))
                        for p in cell.paragraphs:
                            pPr = p._p.get_or_add_pPr()
                            sp = pPr.find(qn("w:spacing"))
                            if sp is not None:
                                sp.set(qn("w:before"), "40")
                                sp.set(qn("w:after"), "40")
                                sp.set(qn("w:line"), "276")
                                sp.set(qn("w:lineRule"), "auto")
                            else:
                                pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="40" w:after="40" w:line="276" w:lineRule="auto"/>'))
                            for r in p.runs:
                                r.font.name = "Times New Roman"
                                r.font.size = Pt(12)
                                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 6. Build Interactive Hyperlinked TOC at End (Chương + Bài học)
    if p_toc_heading:
        print(f"  → Tạo Mục lục tương tác gồm {len(toc_headings)} liên kết...")
        last_elem = p_toc_heading._p

        for item in toc_headings:
            title = html.escape(item["text"])
            bm = item["bm_name"]

            if item.get("is_chapter"):
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:spacing w:before="180" w:after="40"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:b/>
                                <w:sz w:val="23"/>
                                <w:color w:val="0F2A44"/>
                            </w:rPr>
                            <w:t>{title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            elif item.get("is_lesson"):
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:ind w:left="280"/>
                        <w:spacing w:before="20" w:after="20"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:sz w:val="22"/>
                                <w:color w:val="1A4A6B"/>
                            </w:rPr>
                            <w:t>•  {title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            else:
                continue

            p_elem = parse_xml(p_xml)
            last_elem.addnext(p_elem)
            last_elem = p_elem

    # 7. Lưu file docx hoàn thiện
    doc.save(str(docx_path))
    print(f"  ✅ Đã lưu file: {docx_path.name}")

    # 8. Hậu kiểm dọn sạch package XML & ép màu đen 000000 trong styles.xml
    clean_and_blacken_package(docx_path)

def clean_and_blacken_package(docx_path):
    print("🧹 Đang dọn dẹp package XML và ép đen 000000...")
    temp_zip_path = docx_path.parent / f"_temp_clean_{docx_path.name}"
    color_pat = re.compile(r'(<w:color\b[^>]*?\bw:val=[\'"])1E293B([\'"][^>]*?>)', re.IGNORECASE)

    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(temp_zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/styles.xml":
                    text = data.decode("utf-8")
                    text, _ = color_pat.subn(r'\g<1>000000\g<2>', text)
                    text = re.sub(r'w:color\s+w:themeColor="[^"]*"\s+w:themeShade="[^"]*"\s+w:val="000000"', 'w:color w:val="000000"', text)
                    data = text.encode("utf-8")
                elif item.filename == "word/document.xml":
                    text = data.decode("utf-8")
                    text, _ = color_pat.subn(r'\g<1>000000\g<2>', text)
                    data = text.encode("utf-8")
                zout.writestr(item, data)

    temp_zip_path.replace(docx_path)
    print(f"  ✅ Hoàn tất ép đen và dọn sạch XML cho {docx_path.name}")

def main():
    print("================================================================================")
    print("🚀 BẮT ĐẦU XUẤT BẢN SÁCH GIÁO VIÊN C++ QUYỂN 1 (FULL KHUÔN PYTHON)")
    print("================================================================================")

    with open(MANIFEST_FILE, "r", encoding="utf-8") as f:
        manifest_data = json.load(f)

    # 1. Sinh Markdown
    full_md = build_q1_markdown(manifest_data)

    # 2. Convert to DOCX qua Pandoc
    convert_markdown_to_docx(full_md, OUT_DOCX)

    # 3. Post-process Design System (Ảnh minh họa, bảng, màu sắc, TOC)
    post_process_teacher_docx(OUT_DOCX, manifest_data)

    print("\n================================================================================")
    print(f"🎉 XUẤT BẢN THÀNH CÔNG: {OUT_DOCX.name} ({OUT_DOCX.stat().st_size:,} bytes)!")
    print("================================================================================")

if __name__ == "__main__":
    main()
