import re
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

import sys
sys.path.insert(0, '/Users/vu/.gemini/antigravity-ide/brain/3610f3da-8aad-4288-b82c-870385e11baf/scratch')
from generate_l0_problems import PROBLEMS

def clean_latex(text: str) -> str:
    text = re.sub(r'\$([^\$]+)\$', r'\1', text)
    text = text.replace(r'\le', '≤').replace(r'\ge', '≥')
    text = text.replace(r'\ne', '≠').replace(r'\times', '×')
    text = text.replace(r'\dots', '...').replace(r'\text{s}', 's').replace(r'\text{MB}', 'MB')
    text = text.replace(r'\%', '%')
    return text

def format_problem_tag(code: str) -> str:
    parts = code.split('_')
    if len(parts) >= 3:
        return f"{parts[0].upper()}-{parts[1].upper()}-{parts[2]}"
    return code.upper().replace('_', '-')

def set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def create_styled_table(doc, input_lines, output_lines):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.style = 'Table'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Set explicit widths (6800 total -> 3400 each)
    for row in tbl.rows:
        # cantSplit on row
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        
        row.cells[0].width = Inches(3400 / 1440.0)
        row.cells[1].width = Inches(3400 / 1440.0)

    # Header Row
    hdr_titles = ["Đầu vào (Input)", "Đầu ra (Output)"]
    for i, title in enumerate(hdr_titles):
        cell = tbl.rows[0].cells[i]
        set_cell_shading(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        p = cell.paragraphs[0]
        p.style = 'Compact'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Consolas"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Calculate dynamic indent for multiline IO
    def calc_indent_pt(lines):
        max_l = max([len(l) for l in lines] + [1])
        if max_l <= 4:
            return 72.0
        return max(10.0, 72.0 - (max_l - 4) * 3.25)

    in_indent = calc_indent_pt(input_lines)
    out_indent = calc_indent_pt(output_lines)

    for col_idx, (lines, indent_val) in enumerate([(input_lines, in_indent), (output_lines, out_indent)]):
        cell = tbl.rows[1].cells[col_idx]
        set_cell_shading(cell, "FFFFFF")
        set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        for line_idx, line_str in enumerate(lines):
            if line_idx == 0:
                p = cell.paragraphs[0]
            else:
                p = cell.add_paragraph()
            p.style = 'Compact'
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Pt(indent_val)
            
            run = p.add_run(line_str)
            run.font.name = "Consolas"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    return tbl

def build_source_code_p(doc, code_text):
    p = doc.add_paragraph(style='Source Code')
    pPr = p._p.get_or_add_pPr()
    
    # Border & Shading via parse_xml to ensure standard schema
    pbdr = parse_xml(f'''
    <w:pBdr {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="2" w:color="E2E8F0"/>
        <w:left w:val="single" w:sz="8" w:space="4" w:color="CBD5E1"/>
        <w:bottom w:val="single" w:sz="4" w:space="2" w:color="E2E8F0"/>
        <w:right w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
    </w:pBdr>
    ''')
    pPr.append(pbdr)
    
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    pPr.append(shd)
    
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(12.6) # 252 dxa
    
    lines = code_text.strip().splitlines()
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        if idx < len(lines) - 1:
            run.add_break()
            
    return p

print("Loaded helper functions for docx generation successfully.")
