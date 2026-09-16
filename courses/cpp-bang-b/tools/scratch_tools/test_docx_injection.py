import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, Inches, RGBColor
from pathlib import Path

# Load original clean docx
doc = docx.Document('c++-level-1-quyen-1.docx')

# Target paragraph where we want to insert (before Phụ lục B)
target_p = doc.paragraphs[3392]

# Insert a Heading 2 before target_p
h2 = target_p.insert_paragraph_before('Bài tập thực hành', style='Heading 2')
for run in h2.runs:
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

# Insert a Body Text
body_p = target_p.insert_paragraph_before('Dưới đây là 16 bài tập lập trình cơ bản.', style='Body Text')

# Insert a table
tbl = doc.add_table(rows=2, cols=2)
tbl.style = 'Table'
tbl.rows[0].cells[0].text = "Đầu vào (Input)"
tbl.rows[0].cells[1].text = "Đầu ra (Output)"
tbl.rows[1].cells[0].text = "1 2\n3 4"
tbl.rows[1].cells[1].text = "10"

# Move table before target_p in XML body
target_p._p.addprevious(tbl._tbl)

doc.save('test_injected.docx')

# Try reopening
doc2 = docx.Document('test_injected.docx')
print('Successfully saved and reopened test_injected.docx! Paragraph count:', len(doc2.paragraphs))
