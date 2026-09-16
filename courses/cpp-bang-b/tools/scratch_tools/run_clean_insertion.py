import re
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

import sys
sys.path.insert(0, '/Users/vu/.gemini/antigravity-ide/brain/3610f3da-8aad-4288-b82c-870385e11baf/scratch')
from generate_l0_problems import PROBLEMS
from build_clean_l0 import clean_latex, format_problem_tag, create_styled_table, build_source_code_p

DOCX_IN = "c++-level-1-quyen-1.docx"
DOCX_OUT = "c++-level-1-quyen-1.docx"

doc = docx.Document(DOCX_IN)

# 1. Tìm target paragraph cho Phụ lục A (chèn trước Phụ lục B)
target_pb = None
target_c01 = None

for p in doc.paragraphs:
    if p.text.strip() == "Phụ lục B: Lời giải bài tập tham khảo" and p.style.name == "Heading 1":
        target_pb = p
    elif "Chương 01 — Bài 01: Sắp xếp" in p.text and p.style.name == "Heading 2":
        target_c01 = p

if not target_pb:
    raise ValueError("Không tìm thấy Heading 1 Phụ lục B!")
if not target_c01:
    raise ValueError("Không tìm thấy Heading 2 Chương 01 trong Phụ lục B!")

print("Đã xác định vị trí chèn:")
print("- Trước Phụ lục B:", target_pb.text)
print("- Trước Chương 01 trong Phụ lục B:", target_c01.text)

# Helper chèn paragraph hoặc table vào trước một paragraph mục tiêu
def insert_p_before(target_p, text, style='Body Text', bold_prefix=None, space_before=0, space_after=3, line_spacing=13.8):
    new_p = doc.add_paragraph(style=style)
    new_p.paragraph_format.space_before = Pt(space_before)
    new_p.paragraph_format.space_after = Pt(space_after)
    if line_spacing:
        new_p.paragraph_format.line_spacing = Pt(line_spacing)
        
    if bold_prefix:
        r_b = new_p.add_run(bold_prefix)
        r_b.bold = True
        
    if text:
        new_p.add_run(text)
        
    target_p._p.addprevious(new_p._p)
    return new_p

def insert_h3_before(target_p, text):
    new_p = doc.add_paragraph(style='Heading 3')
    new_p.paragraph_format.space_before = Pt(9)
    new_p.paragraph_format.space_after = Pt(3)
    new_p.paragraph_format.line_spacing = Pt(13.8)
    
    run = new_p.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    run.font.size = Pt(13)
    
    target_p._p.addprevious(new_p._p)
    return new_p

# ==========================================
# PHẦN 1: CHÈN 16 BÀI TẬP VÀO PHỤ LỤC A
# ==========================================
# Heading 2 màu đỏ: Bài tập thực hành
h2_bt = doc.add_paragraph(style='Heading 2')
h2_bt.paragraph_format.space_before = Pt(11)
h2_bt.paragraph_format.space_after = Pt(3)
r_h2 = h2_bt.add_run("Bài tập thực hành")
r_h2.bold = True
r_h2.font.name = "Times New Roman"
r_h2.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
r_h2.font.size = Pt(14)
target_pb._p.addprevious(h2_bt._p)

# Intro
insert_p_before(target_pb, "Dưới đây là 16 bài tập lập trình cơ bản bám sát 6 nhóm kỹ năng nền tảng (Biến, Toán tử, Điều kiện, Vòng lặp, Tích lũy, Mảng & Xâu) giúp học sinh rèn luyện tư duy và làm chủ cú pháp C++ chuẩn.", style='BlockText', space_before=3, space_after=6)

for idx, p_info in enumerate(PROBLEMS, 1):
    num_str = f"{idx:02d}"
    code = p_info["code"]
    title = p_info["title"]
    print(f"Adding problem {num_str}: {code}")
    
    tag = format_problem_tag(code)
    h3_text = f"Bài {num_str} [{tag}]: {title}"
    insert_h3_before(target_pb, h3_text)
    
    content = Path(f"problems/{code}/De_Bai.md").read_text(encoding="utf-8")
    def extract_sec(sec_title):
        m = re.search(rf"^## {sec_title}\s*\n(.*?)(?=^## |\Z)", content, re.M | re.S)
        return m.group(1).strip() if m else ""

    boi_canh = clean_latex(extract_sec("Bối cảnh"))
    nhiem_vu = clean_latex(extract_sec("Nhiệm vụ"))
    input_sec = clean_latex(extract_sec("Input"))
    output_sec = clean_latex(extract_sec("Output"))
    sample_sec = extract_sec("Sample 1")
    giai_thich_m = re.search(r"### Giải thích\s*\n(.*?)(?=^## |\Z)", content, re.M | re.S)
    giai_thich = clean_latex(giai_thich_m.group(1).strip() if giai_thich_m else "")
    rang_buoc = clean_latex(extract_sec("Ràng buộc"))

    insert_p_before(target_pb, boi_canh, style='First Paragraph', bold_prefix='Bối cảnh: ')
    insert_p_before(target_pb, nhiem_vu, style='Body Text', bold_prefix='Nhiệm vụ: ')
    
    insert_p_before(target_pb, "", style='Body Text', bold_prefix='Đầu vào (Input):')
    for l in input_sec.splitlines():
        l = l.strip().lstrip('- ').strip()
        if l:
            insert_p_before(target_pb, l, style='Compact', space_before=1, space_after=2)
            
    insert_p_before(target_pb, "", style='First Paragraph', bold_prefix='Đầu ra (Output):')
    for l in output_sec.splitlines():
        l = l.strip().lstrip('- ').strip()
        if l:
            insert_p_before(target_pb, l, style='Compact', space_before=1, space_after=2)
            
    insert_p_before(target_pb, "", style='First Paragraph', bold_prefix='Ví dụ mẫu (Sample 1):')
    in_m = re.search(r"### Input\s*```[a-z]*\n(.*?)```", sample_sec, re.S)
    out_m = re.search(r"### Output\s*```[a-z]*\n(.*?)```", sample_sec, re.S)
    in_lines = in_m.group(1).strip().splitlines() if in_m else [""]
    out_lines = out_m.group(1).strip().splitlines() if out_m else [""]
    tbl = create_styled_table(doc, in_lines, out_lines)
    target_pb._p.addprevious(tbl._tbl)
    
    if giai_thich:
        insert_p_before(target_pb, "", style='Compact', bold_prefix='Giải thích:')
        for l in giai_thich.splitlines():
            l = l.strip()
            if l:
                insert_p_before(target_pb, l, style='Body Text', space_before=1, space_after=2)
                
    insert_p_before(target_pb, "", style='Body Text', bold_prefix='Ràng buộc & Giới hạn:')
    for l in rang_buoc.splitlines():
        l = l.strip().lstrip('- ').strip()
        if l:
            insert_p_before(target_pb, l, style='Compact', space_before=1, space_after=2)

print("Đã chèn toàn bộ 16 bài tập vào Phụ lục A!")

# ==========================================
# PHẦN 2: CHÈN 3 LỜI GIẢI MẪU VÀO PHỤ LỤC B
# ==========================================
h2_pa = doc.add_paragraph(style='Heading 2')
h2_pa.paragraph_format.space_before = Pt(11)
h2_pa.paragraph_format.space_after = Pt(3)
r_h2_pa = h2_pa.add_run("Phụ lục A — Nền tảng C++")
r_h2_pa.bold = True
r_h2_pa.font.name = "Times New Roman"
r_h2_pa.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
r_h2_pa.font.size = Pt(14)
target_c01._p.addprevious(h2_pa._p)

for i in [0, 1, 2]:
    p_info = PROBLEMS[i]
    code = p_info["code"]
    title = p_info["title"]
    code_text = Path(f"problems/{code}/solution.cpp").read_text(encoding="utf-8")
    code_tag = format_problem_tag(code)
    
    h3_sol = doc.add_paragraph(style='Heading 3')
    h3_sol.paragraph_format.space_before = Pt(9)
    h3_sol.paragraph_format.space_after = Pt(3)
    r1 = h3_sol.add_run(code_tag)
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    r1.font.size = Pt(13)
    
    r2 = h3_sol.add_run(f" — {title}")
    r2.bold = True
    r2.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    r2.font.size = Pt(13)
    target_c01._p.addprevious(h3_sol._p)
    
    sc_p = build_source_code_p(doc, code_text)
    target_c01._p.addprevious(sc_p._p)

note_p = doc.add_paragraph()
note_p.paragraph_format.space_before = Pt(8)
note_p.paragraph_format.space_after = Pt(14)
r_note = note_p.add_run("Các bài tập còn lại có phương pháp và cấu trúc tương tự, học sinh tự suy luận và cài đặt.")
r_note.italic = True
r_note.font.size = Pt(11)
target_c01._p.addprevious(note_p._p)

print("Đã chèn 3 lời giải mẫu vào Phụ lục B!")

# Lưu file kết quả
doc.save(DOCX_OUT)
print("✅ Lưu file docx thành công!")

# Thử mở lại bằng python-docx để verify
doc_verify = docx.Document(DOCX_OUT)
print(f"✅ Đã mở lại và xác thực thành công docx: {len(doc_verify.paragraphs)} paragraphs, {len(doc_verify.tables)} tables.")
