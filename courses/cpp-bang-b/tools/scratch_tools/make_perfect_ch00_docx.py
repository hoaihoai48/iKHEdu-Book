import re
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import xml.etree.ElementTree as ET

import sys
sys.path.insert(0, '/Users/vu/.gemini/antigravity-ide/brain/3610f3da-8aad-4288-b82c-870385e11baf/scratch')
from generate_l0_problems import PROBLEMS

DOCX = "c++-level-1-quyen-1.docx"
doc = docx.Document(DOCX)
body_elem = doc._body._element

# 1. Tìm vị trí khối Chương 00 hiện tại: từ P11 đến ngay trước P_CH01
h1_ch00 = None
h1_ch01 = None

for p in doc.paragraphs:
    t = p.text.strip()
    if p.style.name == "Heading 1" and "CHƯƠNG 00" in t:
        h1_ch00 = p
    elif p.style.name == "Heading 1" and "CHƯƠNG 01" in t:
        h1_ch01 = p
        break

if not h1_ch00 or not h1_ch01:
    raise ValueError("Không tìm thấy Heading Chương 00 hoặc Chương 01!")

all_children = list(body_elem)
start_idx = all_children.index(h1_ch00._p)
end_idx = all_children.index(h1_ch01._p)

print(f"Xóa toàn bộ khối Chương 00 cũ từ index {start_idx} đến {end_idx} (số phần tử: {end_idx - start_idx})")
old_ch00_elems = all_children[start_idx:end_idx]
for elem in old_ch00_elems:
    body_elem.remove(elem)

# Helper làm sạch text công thức toán và lệnh LaTeX
def clean_latex(text: str) -> str:
    # 1. Phân số: \frac{a}{b} -> a / b
    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'\1 / \2', text)
    # 2. Xóa các ký hiệu toán phổ biến
    text = text.replace(r'\le', '≤').replace(r'\ge', '≥')
    text = text.replace(r'\ne', '≠').replace(r'\times', '×')
    text = text.replace(r'\dots', '...').replace(r'\cdots', '...')
    text = text.replace(r'\%', '%')
    text = text.replace(r'\approx', '≈').replace(r'\infty', '∞')
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\mathcal\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', text)
    # 3. Bỏ dấu $ bao quanh công thức đơn
    text = re.sub(r'\$([^\$]+)\$', r'\1', text)
    return text

def format_problem_tag(code: str) -> str:
    parts = code.split('_')
    if len(parts) >= 3:
        return f"{parts[0].upper()}-{parts[1].upper()}-{parts[2]}"
    return code.upper().replace('_', '-')

# Helper tạo paragraph Bullet (Compact + numPr numId=9)
def create_bullet_p(text: str):
    p = doc.add_paragraph(style='Compact')
    pPr = p._p.get_or_add_pPr()
    numPr = parse_xml(f'''
    <w:numPr {nsdecls("w")}>
        <w:ilvl w:val="0"/>
        <w:numId w:val="9"/>
    </w:numPr>
    ''')
    pPr.append(numPr)
    
    # spacing & jc
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing')
        pPr.append(sp)
    sp.set(qn('w:line'), '276')
    sp.set(qn('w:lineRule'), 'auto')
    
    jc = pPr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc')
        pPr.append(jc)
    jc.set(qn('w:val'), 'both')
    
    run = p.add_run(clean_latex(text))
    return p

# Helper tạo sample table (y hệt Table 18 của SX-01) với tblBorders rõ nét
def create_sample_table(input_lines, output_lines):
    def calc_indent(lines):
        max_l = max([len(l) for l in lines] + [1])
        if max_l <= 4:
            return 1440
        return max(200, int(1440 - (max_l - 4) * 65))

    in_indent = calc_indent(input_lines)
    out_indent = calc_indent(output_lines)
    
    in_ps = ""
    for l in input_lines:
        in_ps += f'''
        <w:p>
            <w:pPr>
                <w:pStyle w:val="Compact"/>
                <w:spacing w:before="40" w:after="40"/>
                <w:ind w:left="{in_indent}"/>
            </w:pPr>
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>
                    <w:color w:val="0F172A"/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:t xml:space="preserve">{l}</w:t>
            </w:r>
        </w:p>
        '''
        
    out_ps = ""
    for l in output_lines:
        out_ps += f'''
        <w:p>
            <w:pPr>
                <w:pStyle w:val="Compact"/>
                <w:spacing w:before="40" w:after="40"/>
                <w:ind w:left="{out_indent}"/>
            </w:pPr>
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>
                    <w:color w:val="0F172A"/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:t xml:space="preserve">{l}</w:t>
            </w:r>
        </w:p>
        '''

    tbl_xml = f'''
    <w:tbl {nsdecls("w")}>
        <w:tblPr>
            <w:tblStyle w:val="Table"/>
            <w:tblW w:type="dxa" w:w="6800"/>
            <w:jc w:val="center"/>
            <w:tblBorders>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            </w:tblBorders>
            <w:tblLook w:firstColumn="1" w:firstRow="1" w:lastColumn="0" w:lastRow="0" w:noHBand="0" w:noVBand="1" w:val="04A0"/>
        </w:tblPr>
        <w:tblGrid>
            <w:gridCol w:w="4949"/>
            <w:gridCol w:w="4949"/>
        </w:tblGrid>
        <w:tr>
            <w:trPr><w:cantSplit/></w:trPr>
            <w:tc>
                <w:tcPr>
                    <w:tcW w:type="dxa" w:w="3400"/>
                    <w:shd w:fill="F1F5F9"/>
                    <w:tcMar>
                        <w:top w:type="dxa" w:w="60"/>
                        <w:bottom w:type="dxa" w:w="60"/>
                        <w:left w:type="dxa" w:w="60"/>
                        <w:right w:type="dxa" w:w="60"/>
                    </w:tcMar>
                    <w:vAlign w:val="center"/>
                </w:tcPr>
                <w:p>
                    <w:pPr>
                        <w:pStyle w:val="Compact"/>
                        <w:spacing w:before="60" w:after="60"/>
                        <w:jc w:val="center"/>
                    </w:pPr>
                    <w:r>
                        <w:rPr>
                            <w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>
                            <w:b/>
                            <w:color w:val="1E293B"/>
                            <w:sz w:val="22"/>
                        </w:rPr>
                        <w:t>Đầu vào (Input)</w:t>
                    </w:r>
                </w:p>
            </w:tc>
            <w:tc>
                <w:tcPr>
                    <w:tcW w:type="dxa" w:w="3400"/>
                    <w:shd w:fill="F1F5F9"/>
                    <w:tcMar>
                        <w:top w:type="dxa" w:w="60"/>
                        <w:bottom w:type="dxa" w:w="60"/>
                        <w:left w:type="dxa" w:w="60"/>
                        <w:right w:type="dxa" w:w="60"/>
                    </w:tcMar>
                    <w:vAlign w:val="center"/>
                </w:tcPr>
                <w:p>
                    <w:pPr>
                        <w:pStyle w:val="Compact"/>
                        <w:spacing w:before="60" w:after="60"/>
                        <w:jc w:val="center"/>
                    </w:pPr>
                    <w:r>
                        <w:rPr>
                            <w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>
                            <w:b/>
                            <w:color w:val="1E293B"/>
                            <w:sz w:val="22"/>
                        </w:rPr>
                        <w:t>Đầu ra (Output)</w:t>
                    </w:r>
                </w:p>
            </w:tc>
        </w:tr>
        <w:tr>
            <w:trPr><w:cantSplit/></w:trPr>
            <w:tc>
                <w:tcPr>
                    <w:tcW w:type="dxa" w:w="3400"/>
                    <w:shd w:fill="FFFFFF"/>
                    <w:tcMar>
                        <w:top w:type="dxa" w:w="40"/>
                        <w:bottom w:type="dxa" w:w="40"/>
                        <w:left w:type="dxa" w:w="60"/>
                        <w:right w:type="dxa" w:w="60"/>
                    </w:tcMar>
                    <w:vAlign w:val="center"/>
                </w:tcPr>
                {in_ps}
            </w:tc>
            <w:tc>
                <w:tcPr>
                    <w:tcW w:type="dxa" w:w="3400"/>
                    <w:shd w:fill="FFFFFF"/>
                    <w:tcMar>
                        <w:top w:type="dxa" w:w="40"/>
                        <w:bottom w:type="dxa" w:w="40"/>
                        <w:left w:type="dxa" w:w="60"/>
                        <w:right w:type="dxa" w:w="60"/>
                    </w:tcMar>
                    <w:vAlign w:val="center"/>
                </w:tcPr>
                {out_ps}
            </w:tc>
        </w:tr>
    </w:tbl>
    '''
    return parse_xml(tbl_xml)

# Tạo danh sách các phần tử của Chương 00 mới
new_ch00_elems = []

# 1. Heading 1 Chương 00
p_h1_ch00 = doc.add_paragraph()
p_h1_ch00.style = "Heading 1"
pPr = p_h1_ch00._p.get_or_add_pPr()
pPr.append(OxmlElement('w:pageBreakBefore'))
sp = OxmlElement('w:spacing')
sp.set(qn('w:after'), '100')
pPr.append(sp)
jc = OxmlElement('w:jc')
jc.set(qn('w:val'), 'center')
pPr.append(jc)

r1 = p_h1_ch00.add_run("CHƯƠNG 00")
r1.bold = True
r1.font.name = "Times New Roman"
r1.font.size = Pt(18)
r1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

r_br = p_h1_ch00.add_run()
r_br.add_break()

r2 = p_h1_ch00.add_run("ÔN TẬP & NỀN TẢNG LẬP TRÌNH C++")
r2.bold = True
r2.font.name = "Times New Roman"
r2.font.size = Pt(18)
r2.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

new_ch00_elems.append(p_h1_ch00._p)

# 2. Heading 1 Bài 00
p_b00 = doc.add_paragraph()
p_b00.style = "Heading 1"
pPr = p_b00._p.get_or_add_pPr()
sp = OxmlElement('w:spacing')
sp.set(qn('w:before'), '200')
sp.set(qn('w:after'), '60')
pPr.append(sp)

r = p_b00.add_run("Bài 00: Ôn tập kỹ năng lập trình cơ bản")
r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(15.5)
r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
r.font.underline = True

new_ch00_elems.append(p_b00._p)

# 3. Lời dẫn Block Text
p_intro = doc.add_paragraph(style='Block Text')
p_intro.paragraph_format.space_before = Pt(4)
p_intro.paragraph_format.space_after = Pt(10)
p_intro.add_run("Dưới đây là 16 bài tập lập trình cơ bản bám sát 6 nhóm kỹ năng nền tảng (Biến, Toán tử, Điều kiện, Vòng lặp, Tích lũy, Mảng & Xâu) giúp học sinh rèn luyện tư duy và làm chủ cú pháp C++ chuẩn.")

new_ch00_elems.append(p_intro._p)

# 4. Heading 2 màu đỏ: Bài tập thực hành
p_bt_red = doc.add_paragraph(style='Heading 2')
pPr = p_bt_red._p.get_or_add_pPr()
sp = OxmlElement('w:spacing')
sp.set(qn('w:before'), '220')
sp.set(qn('w:after'), '60')
pPr.append(sp)
rPr_p = OxmlElement('w:rPr')
rPr_p.append(OxmlElement('w:b'))
col = OxmlElement('w:color')
col.set(qn('w:val'), 'FF0000')
rPr_p.append(col)
pPr.append(rPr_p)

r_red = p_bt_red.add_run("Bài tập thực hành")
r_red.bold = True
r_red.font.name = "Times New Roman"
r_red.font.size = Pt(14)
r_red.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

new_ch00_elems.append(p_bt_red._p)

# 5. Duyệt qua 16 bài toán L0
for idx, p_info in enumerate(PROBLEMS, 1):
    num_str = f"{idx:02d}"
    code = p_info["code"]
    title = p_info["title"]
    tag = format_problem_tag(code)
    
    # 5.1 Heading 3: Bài XX [CPPB-L0-XX]: Tên Bài
    p_h3 = doc.add_paragraph(style='Heading 3')
    pPr = p_h3._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '180')
    sp.set(qn('w:after'), '60')
    sp.set(qn('w:line'), '276')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    
    r_h3 = p_h3.add_run(f"Bài {num_str} [{tag}]: {title}")
    r_h3.bold = True
    r_h3.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    r_h3.font.size = Pt(13)
    new_ch00_elems.append(p_h3._p)
    
    # Đọc De_Bai.md
    content = Path(f"problems/{code}/De_Bai.md").read_text(encoding="utf-8")
    def extract_sec(sec_title):
        m = re.search(rf"^## {sec_title}\s*\n(.*?)(?=^## |\Z)", content, re.M | re.S)
        return m.group(1).strip() if m else ""

    boi_canh = clean_latex(extract_sec("Bối cảnh"))
    nhiem_vu = clean_latex(extract_sec("Nhiệm vụ"))
    input_sec = extract_sec("Input")
    output_sec = extract_sec("Output")
    sample_sec = extract_sec("Sample 1")
    giai_thich_m = re.search(r"### Giải thích\s*\n(.*?)(?=^## |\Z)", content, re.M | re.S)
    giai_thich = clean_latex(giai_thich_m.group(1).strip() if giai_thich_m else "")
    rang_buoc = extract_sec("Ràng buộc")
    
    # 5.2 Bối cảnh (First Paragraph)
    p_bc = doc.add_paragraph(style='First Paragraph')
    r_bc_lbl = p_bc.add_run("Bối cảnh:")
    r_bc_lbl.bold = True
    r_bc_txt = p_bc.add_run(f" {boi_canh}")
    new_ch00_elems.append(p_bc._p)
    
    # 5.3 Nhiệm vụ (Body Text)
    p_nv = doc.add_paragraph(style='Body Text')
    r_nv_lbl = p_nv.add_run("Nhiệm vụ:")
    r_nv_lbl.bold = True
    r_nv_txt = p_nv.add_run(f" {nhiem_vu}")
    new_ch00_elems.append(p_nv._p)
    
    # 5.4 Đầu vào (Input): (Body Text)
    p_in_lbl = doc.add_paragraph(style='Body Text')
    r_in_lbl = p_in_lbl.add_run("Đầu vào (Input):")
    r_in_lbl.bold = True
    new_ch00_elems.append(p_in_lbl._p)
    
    # Bullet input
    for l in input_sec.splitlines():
        l = l.strip().lstrip('- ').strip()
        if l:
            p_in_b = create_bullet_p(l)
            new_ch00_elems.append(p_in_b._p)
            
    # 5.5 Đầu ra (Output): (First Paragraph)
    p_out_lbl = doc.add_paragraph(style='First Paragraph')
    r_out_lbl = p_out_lbl.add_run("Đầu ra (Output):")
    r_out_lbl.bold = True
    new_ch00_elems.append(p_out_lbl._p)
    
    # Bullet output
    for l in output_sec.splitlines():
        l = l.strip().lstrip('- ').strip()
        if l:
            p_out_b = create_bullet_p(l)
            new_ch00_elems.append(p_out_b._p)
            
    # 5.6 Ví dụ mẫu (Sample 1): (First Paragraph)
    p_sm_lbl = doc.add_paragraph(style='First Paragraph')
    r_sm_lbl = p_sm_lbl.add_run("Ví dụ mẫu (Sample 1):")
    r_sm_lbl.bold = True
    new_ch00_elems.append(p_sm_lbl._p)
    
    # Table Sample
    in_m = re.search(r"### Input\s*```[a-z]*\n(.*?)```", sample_sec, re.S)
    out_m = re.search(r"### Output\s*```[a-z]*\n(.*?)```", sample_sec, re.S)
    in_lines = in_m.group(1).strip().splitlines() if in_m else [""]
    out_lines = out_m.group(1).strip().splitlines() if out_m else [""]
    tbl_elem = create_sample_table(in_lines, out_lines)
    new_ch00_elems.append(tbl_elem)
    
    # 5.7 Giải thích: (Compact bold)
    if giai_thich:
        p_gt_lbl = doc.add_paragraph(style='Compact')
        r_gt_lbl = p_gt_lbl.add_run("Giải thích:")
        r_gt_lbl.bold = True
        new_ch00_elems.append(p_gt_lbl._p)
        
        # Thân giải thích dùng Normal (như P322 của SX-01)
        for l in giai_thich.splitlines():
            l = l.strip().lstrip('- ').strip()
            if l:
                p_gt = doc.add_paragraph(style='Normal')
                p_gt.add_run(clean_latex(l))
                new_ch00_elems.append(p_gt._p)
                
    # 5.8 Ràng buộc & Giới hạn: (Body Text)
    p_rb_lbl = doc.add_paragraph(style='Body Text')
    r_rb_lbl = p_rb_lbl.add_run("Ràng buộc & Giới hạn:")
    r_rb_lbl.bold = True
    new_ch00_elems.append(p_rb_lbl._p)
    
    # Bullet ràng buộc
    for l in rang_buoc.splitlines():
        l = l.strip().lstrip('- ').strip()
        if l:
            p_rb_b = create_bullet_p(l)
            new_ch00_elems.append(p_rb_b._p)

print(f"Tổng số elements tạo mới cho Chương 00: {len(new_ch00_elems)}")

# Chèn toàn bộ new_ch00_elems vào trước h1_ch01
for elem in new_ch00_elems:
    h1_ch01._p.addprevious(elem)

print("Đã chèn toàn bộ Chương 00 chuẩn format 100% vào trước Chương 01!")

# Lưu file
doc.save(DOCX)
print("✅ Đã lưu file docx hoàn tất!")
