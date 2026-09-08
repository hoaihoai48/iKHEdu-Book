#!/usr/bin/env python3
"""
BUILD SCRIPT: SÁCH GIÁO VIÊN C++ TẬP 3 (BÀI 15 - 21, 105 BÀI TOÁN)
Chuẩn khuôn 100% theo build_word_teacher_q1.py / q2.py:
1. Hình ảnh minh họa trực quan bối cảnh (Pure Scenario Illustration) gắn ngay sau Bối cảnh của từng bài.
2. Tiêu đề ảnh italic (Times New Roman 10pt, màu #475569).
3. Bảng Sample IO: Header nền #F1F5F9, chữ Consolas 11pt Bold màu #000000, data cell nền trắng #FFFFFF với Dynamic Left Indent (72pt -> 10pt).
4. Đệm cách trước và sau mọi bảng: paragraph trước w:after="120" (6pt), paragraph sau w:before="120" (6pt).
5. Khung Code: Consolas 9pt, line 1.05, nền #F8FAFC, viền #E2E8F0, căn trái 100%, 0 numPr.
6. Toàn bộ chữ Body và Headings dùng màu đen tuyền #000000 (độ tương phản sắc nét in ấn).
7. Header trang: Tab Stop right w:pos="9899", Watermark logo image13.jpeg đầy đủ 3 headers, đường kẻ ngang Straight Connector 1 màu đen 1.5pt.
   Header text: "Giáo trình C++ - Sách giáo viên - Quyển 3"
8. Mục lục: Hyperlink tương tác các Chương + Bài học vừa vặn 1 trang, không vỡ trang.
"""

import os
import re
import sys
import json
import html
import shutil
import zipfile
import subprocess
import lxml.etree as ET
from pathlib import Path

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path("/Users/vu/Developer/ikhEdu_lessons/courses/cpp-bang-b")
PROBLEMS_DIR = BASE_DIR / "problems"
MANIFEST_FILE = BASE_DIR / "word_build_manifest_gv.json"
PREFACE_FILE = BASE_DIR / "reference" / "Loi_Noi_Dau_GV.md"
TEMPLATE_DOCX = BASE_DIR / "c++-level-1-quyen-1.docx"
OUT_DOCX = BASE_DIR / "cpp-giaovien-quyen-3.docx"
ASSETS_PNG_DIR = BASE_DIR / "assets_png"

CHAPTER_NAMES = {
    1: "CHƯƠNG 01: THUẬT TOÁN SẮP XẾP & KỸ THUẬT MẢNG",
    2: "CHƯƠNG 02: MẢNG TIỀN TỐ & TÌM KIẾM NHỊ PHÂN",
    3: "CHƯƠNG 03: SỐ HỌC & ĐẠI SỐ MODULAR",
    4: "CHƯƠNG 04: ĐỆ QUY, CHIA ĐỂ TRỊ & QUAY LUI",
    5: "CHƯƠNG 05: QUY HOẠCH ĐỘNG (DYNAMIC PROGRAMMING)",
    6: "CHƯƠNG 06: CẤU TRÚC DỮ LIỆU NÂNG CAO",
    7: "CHƯƠNG 07: ĐỒ THỊ & CÂY TRUY VẤN ĐOẠN",
}

# Load problem illustrations map (combine q1 and q2 maps)
DIAGRAMS_MAP = {}
for map_name in ["problem_diagrams_q1_map.json", "problem_diagrams_q2_map.json"]:
    map_file = BASE_DIR / map_name
    if map_file.exists():
        with open(map_file, "r", encoding="utf-8") as f:
            DIAGRAMS_MAP.update(json.load(f))

def parse_de_bai(de_bai_path):
    if not de_bai_path.exists():
        return {"bcanh": "", "nv": "", "inp": "", "out": "", "samples": []}
    text = de_bai_path.read_text(encoding="utf-8")

    bcanh_m = re.search(r'## Bối cảnh\s*\n(.*?)(?=\n## Nhiệm vụ|\n## Input|\Z)', text, re.DOTALL)
    nv_m = re.search(r'## Nhiệm vụ\s*\n(.*?)(?=\n## Input|\n## Output|\Z)', text, re.DOTALL)
    inp_m = re.search(r'## Input\s*\n(.*?)(?=\n## Output|\n## Sample|\Z)', text, re.DOTALL)
    out_m = re.search(r'## Output\s*\n(.*?)(?=\n## Sample|\n## Ràng buộc|\Z)', text, re.DOTALL)

    bcanh = bcanh_m.group(1).strip() if bcanh_m else ""
    nv = nv_m.group(1).strip() if nv_m else ""
    inp = inp_m.group(1).strip() if inp_m else ""
    out = out_m.group(1).strip() if out_m else ""

    samples = []
    sample_blocks = re.findall(r'## (Sample \d+)\s*\n(.*?)(?=\n## Sample \d+|\n## Ràng buộc|\Z)', text, re.DOTALL)
    for s_title, s_body in sample_blocks:
        sinp_m = re.search(r'### Input\s*```(?:text)?\n(.*?)```', s_body, re.DOTALL)
        sout_m = re.search(r'### Output\s*```(?:text)?\n(.*?)```', s_body, re.DOTALL)
        sexpl_m = re.search(r'### Giải thích\s*\n(.*)', s_body, re.DOTALL)

        sinp = sinp_m.group(1).strip() if sinp_m else ""
        sout = sout_m.group(1).strip() if sout_m else ""
        sexpl = sexpl_m.group(1).strip() if sexpl_m else ""

        samples.append({
            "title": s_title,
            "input": sinp,
            "output": sout,
            "explain": sexpl
        })

    return {
        "bcanh": bcanh,
        "nv": nv,
        "inp": inp,
        "out": out,
        "samples": samples
    }

def parse_guide(guide_path):
    if not guide_path.exists():
        return []
    text = guide_path.read_text(encoding="utf-8")

    sections = []
    raw_secs = re.split(r'\n(?=##\s+\d+\.)', text)

    for raw in raw_secs:
        m = re.match(r'##\s+(\d+\.\s+[^\n]+)\n(.*)', raw.strip(), re.DOTALL)
        if m:
            sec_title = m.group(1).strip()
            sec_body = m.group(2).strip()
            sec_body = re.sub(r'\n---\s*$', '', sec_body).strip()
            sec_body = re.sub(r'^(#{1,3})\s+', '##### ', sec_body, flags=re.MULTILINE)
            sections.append({
                "title": sec_title,
                "body": sec_body
            })

    return sections

def format_teacher_problem_markdown(prob_idx, prob_code, prob_title, parsed_de_bai, sol_code, parsed_guide):
    md = []
    md.append(f"### Bài {prob_idx:02d} [{prob_code}]: {prob_title}\n\n")

    if parsed_de_bai["bcanh"]:
        md.append(f"Bối cảnh: {parsed_de_bai['bcanh']}\n\n")

    if parsed_de_bai["nv"]:
        md.append(f"Nhiệm vụ: {parsed_de_bai['nv']}\n\n")

    if parsed_de_bai["inp"]:
        md.append(f"**Đầu vào (Input):**\n\n{parsed_de_bai['inp']}\n\n")

    if parsed_de_bai["out"]:
        md.append(f"**Đầu ra (Output):**\n\n{parsed_de_bai['out']}\n\n")

    for s in parsed_de_bai["samples"]:
        s_title_str = f"Ví dụ mẫu ({s['title']}):"
        inp_str = s["input"].replace("\n", " <br> ")
        out_str = s["output"].replace("\n", " <br> ")

        md.append(f"**{s_title_str}**\n\n")
        md.append("| Đầu vào (Input) | Đầu ra (Output) |\n")
        md.append("| :--- | :--- |\n")
        md.append(f"| {inp_str} | {out_str} |\n\n")

        if s["explain"]:
            md.append(f"**Giải thích:** {s['explain']}\n\n")

    for g_sec in parsed_guide:
        if "Lời giải tham khảo" in g_sec["title"]:
            continue
        md.append(f"#### {g_sec['title']}\n\n")
        md.append(f"{g_sec['body']}\n\n")

    md.append("#### 4. Lời giải tham khảo\n\n")
    md.append(f"```cpp\n{sol_code}\n```\n\n")

    return "".join(md)

def build_tap3_markdown(manifest_data):
    md_parts = []
    
    # Lời nói đầu (Heading 1)
    md_parts.append("# Lời nói đầu\n\n")
    md_parts.append(
        "Cuốn **Giáo trình C++ — Sách giáo viên — Quyển 3** là tập hoàn thiện bộ tài liệu đào tạo "
        "lập trình thuật toán C++ chuyên sâu iKHEDU, dành riêng cho quý thầy cô giáo và các huấn luyện viên chuyên trách đội tuyển tin học. "
        "Tài liệu này cung cấp trọn bộ lời giải chi tiết, mã nguồn C++ tối ưu và hệ thống **phương pháp luận giảng dạy thực chiến chuyên sâu** "
        "(ý tưởng và phân tích thuật toán, bảng mô phỏng chạy tay từng bước trên số liệu mẫu thực tế, lưu ý và các bẫy lỗi kinh điển thường gặp) "
        "cho toàn bộ **105 bài toán thực hành** từ Bài 15 đến Bài 21 thuộc 3 trụ cột thuật toán đỉnh cao: Quy hoạch động chuỗi (LCS & Edit Distance), "
        "Cấu trúc dữ liệu STL nâng cao (Set, Map, Heap), Ngăn xếp đơn điệu (Monotonic Stack), Hàng đợi hai đầu (Monotonic Deque), "
        "Lý thuyết đồ thị cơ bản (Duyệt BFS & DFS), Đồ thị lưới 2D & Flood Fill, cùng Cấu trúc cây phân đoạn (Segment Tree) & Fenwick Tree (BIT).\n\n"
        "Mỗi bài toán trong sách giáo viên được cấu trúc thống nhất và tinh gọn qua 4 mục trọng tâm: "
        "(1) Ý tưởng & Phân tích thuật toán gắn liền với bản chất toán học; "
        "(2) Bảng chạy tay trực quan (Dry Run Table) trên các giá trị số thực tế của mẫu thử (Sample); "
        "(3) Lưu ý & các bẫy lỗi kinh điển học sinh thường gặp trong môi trường lập trình; và "
        "(4) Mã nguồn C++ tham khảo tối ưu nhất (Fast I/O, Safe Input, 0 std::). "
        "Tài liệu là cẩm nang đắc lực hỗ trợ thầy cô tổ chức các giờ dạy chất lượng cao, phân tích sai sót của học sinh và đồng hành hiệu quả trong từng bài học.\n\n"
    )

    current_chapter = None
    tap3_lessons = manifest_data["lessons"][14:21]
    prob_count = 0

    for lesson_info in tap3_lessons:
        chap_num = lesson_info["chapter"]
        chap_title = CHAPTER_NAMES.get(chap_num, f"CHƯƠNG {chap_num:02d}")
        lesson_title = lesson_info["title"]

        if chap_num != current_chapter:
            current_chapter = chap_num
            md_parts.append(f"# {chap_title}\n\n")

        md_parts.append(f"# {lesson_title}\n\n")

        for p_idx, prob in enumerate(lesson_info["problems"], 1):
            prob_code = prob["code"]
            prob_title = prob["title"]
            de_bai_file = BASE_DIR / prob["de_bai"]
            sol_file = BASE_DIR / prob["solution"]
            guide_file = BASE_DIR / prob["guide"]

            parsed_de_bai = parse_de_bai(de_bai_file)
            sol_code = sol_file.read_text(encoding="utf-8").strip() if sol_file.exists() else ""
            parsed_guide = parse_guide(guide_file)

            p_md = format_teacher_problem_markdown(p_idx, prob_code, prob_title, parsed_de_bai, sol_code, parsed_guide)
            md_parts.append(p_md)
            prob_count += 1

    md_parts.append("# Mục lục\n\n")
    print(f"  → Tổng hợp Markdown Quyển 3: {len(tap3_lessons)} bài học, {prob_count} bài toán.")
    return "".join(md_parts)

def convert_markdown_to_docx(full_md, out_docx):
    print("🔧 Đang chạy Pandoc Markdown → DOCX...")
    intermediate_md = BASE_DIR / "_build_intermediate_gv_tap3.md"
    intermediate_md.write_text(full_md, encoding="utf-8")

    cmd = [
        "pandoc",
        str(intermediate_md),
        "-o", str(out_docx),
        "--reference-doc", str(TEMPLATE_DOCX),
        "--from=markdown+tex_math_dollars+pipe_tables+fenced_code_blocks+yaml_metadata_block+header_attributes",
        "--to=docx",
        "--standalone",
        "--wrap=auto",
        "--columns=80"
    ]

    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"❌ Pandoc lỗi: {res.stderr}")
        sys.exit(1)

    print(f"  ✅ Pandoc thành công: {out_docx.name} ({out_docx.stat().st_size:,} bytes)")
    if intermediate_md.exists():
        intermediate_md.unlink()
    return True

def insert_picture_with_caption(doc, img_filename, caption_text, target_p_elem):
    """Chèn hình minh họa bối cảnh centered và chú thích italic vào trước target_p_elem."""
    img_path = ASSETS_PNG_DIR / img_filename
    if not img_path.exists():
        return

    # 1. Paragraph chứa hình ảnh
    p_pic = doc.add_paragraph()
    p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pic.paragraph_format.space_before = Pt(6)
    p_pic.paragraph_format.space_after = Pt(2)
    p_pic.paragraph_format.line_spacing = 1.0
    r_pic = p_pic.add_run()
    r_pic.add_picture(str(img_path), width=Inches(5.0))
    target_p_elem.addprevious(p_pic._p)

    # 2. Paragraph chứa caption
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(6)
    p_cap.paragraph_format.line_spacing = 1.0
    r_cap = p_cap.add_run(f"Hình minh họa: {caption_text}")
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    target_p_elem.addprevious(p_cap._p)

def post_process_teacher_docx(docx_path, manifest_data):
    print("🎨 Đang hậu xử lý Design System chuẩn sách giáo viên Quyển 3...")
    doc = docx.Document(str(docx_path))
    s = doc.sections[0]

    # 1. Khổ giấy & Lề sách
    s.page_width = Pt(595.3)   # A4 Width
    s.page_height = Pt(841.9)  # A4 Height
    s.top_margin = Pt(36.0)
    s.bottom_margin = Pt(36.0)
    s.left_margin = Pt(64.35)  # Gáy sách
    s.right_margin = Pt(36.0)

    # 2. Update Headers (3 headers: even, default, first)
    header_book_text = "Giáo trình C++ - Sách giáo viên - Quyển 3"
    for h in [s.even_page_header, s.header, s.first_page_header]:
        p0 = h.paragraphs[0]
        pPr = p0._p.get_or_add_pPr()
        for tb in pPr.findall(qn("w:tabs")): pPr.remove(tb)
        pPr.append(parse_xml(f'<w:tabs {nsdecls("w")}><w:tab {nsdecls("w")} w:val="right" w:pos="9899"/></w:tabs>'))

        watermark_runs = []
        for r in p0.runs:
            if r._r.find(qn("w:pict")) is not None:
                watermark_runs.append(r._r)

        for child in list(p0._p):
            if child.tag.split("}")[-1] != "pPr":
                p0._p.remove(child)

        for wr in watermark_runs:
            p0._p.append(wr)

        r_left = parse_xml(f'''<w:r {nsdecls("w")}>
            <w:rPr>
                <w:rFonts w:cs="Times New Roman"/>
                <w:b/><w:bCs/><w:i/><w:iCs/>
                <w:sz w:val="22"/><w:szCs w:val="22"/>
            </w:rPr>
            <w:t>Trung tâm giáo dục &amp; đào tạo tin học iKH</w:t>
        </w:r>''')
        p0._p.append(r_left)

        r_tab = parse_xml(f'''<w:r {nsdecls("w")}>
            <w:rPr><w:rFonts w:cs="Times New Roman"/><w:sz w:val="22"/><w:szCs w:val="22"/></w:rPr>
            <w:tab/>
        </w:r>''')
        p0._p.append(r_tab)

        r_right = parse_xml(f'''<w:r {nsdecls("w")}>
            <w:rPr>
                <w:rFonts w:cs="Times New Roman"/>
                <w:i/><w:iCs/><w:noProof/>
                <w:sz w:val="22"/><w:szCs w:val="22"/>
                <w:lang w:val="vi-VN"/>
            </w:rPr>
            <w:t>{html.escape(header_book_text)}</w:t>
        </w:r>''')
        p0._p.append(r_right)

    # 3. Process Paragraphs, Headings & Bookmarks
    toc_headings = []
    p_toc_heading = None

    for idx, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else ""
        text = para.text.strip()

        # Heading 1
        if style_name == "Heading 1" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading1"):
            para.paragraph_format.keep_with_next = True
            is_chapter = "CHƯƠNG" in text.upper()
            is_preface = "LỜI NÓI ĐẦU" in text.upper()
            is_toc = "MỤC LỤC" in text.upper()

            if is_toc:
                p_toc_heading = para
                para.paragraph_format.page_break_before = True
                continue

            if is_preface:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(14)
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(14)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            elif is_chapter:
                para.paragraph_format.page_break_before = True
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before = Pt(360 / 20)  # 18pt
                para.paragraph_format.space_after = Pt(100 / 20)   # 5pt
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(16)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            else:
                # Lesson heading
                # Nếu là bài học ngay sau tiêu đề Chương thì KHÔNG ngắt trang (cùng nằm trên 1 trang với Chương)
                prev_text = doc.paragraphs[idx - 1].text.strip() if idx > 0 else ""
                if prev_text.startswith("CHƯƠNG"):
                    para.paragraph_format.page_break_before = False
                    # Xóa triệt để thẻ w:pageBreakBefore để Word không ngắt trang
                    if para._p.pPr is not None:
                        for pbb_el in para._p.pPr.findall(qn("w:pageBreakBefore")):
                            para._p.pPr.remove(pbb_el)
                else:
                    para.paragraph_format.page_break_before = True
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.space_before = Pt(14)
                para.paragraph_format.space_after = Pt(6)
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(15)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

            # Bookmark cho TOC
            if not is_preface and not is_toc:
                bm_name = f"_toc_bm_t3_{len(toc_headings)}"
                bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{len(toc_headings)}" w:name="{bm_name}"/>')
                bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{len(toc_headings)}"/>')
                para._p.insert(0, bm_start)
                para._p.append(bm_end)

                toc_headings.append({
                    "text": text,
                    "bm_name": bm_name,
                    "is_chapter": is_chapter,
                    "is_lesson": not is_chapter and not is_preface and not is_toc
                })

        # Heading 2 ("Bài tập thực hành")
        elif style_name == "Heading 2" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading2"):
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(12)
            para.paragraph_format.space_after = Pt(4)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14)
                r.font.bold = True
                if "BÀI TẬP THỰC HÀNH" in text.upper():
                    r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
                else:
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Heading 3 (Tên bài toán)
        elif style_name == "Heading 3" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading3"):
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(14)
            para.paragraph_format.space_after = Pt(4)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Heading 4 (Mục sư phạm)
        elif style_name == "Heading 4" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading4"):
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(10)
            para.paragraph_format.space_after = Pt(3)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Heading 5 (Tiêu đề con trong hướng dẫn)
        elif style_name == "Heading 5" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "Heading5"):
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(2)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Source Code
        elif style_name == "Source Code" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "SourceCode"):
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pPr = para._p.get_or_add_pPr()
            for numPr in pPr.findall(qn("w:numPr")): pPr.remove(numPr)
            for r in para.runs:
                r.font.name = "Consolas"
                r.font.size = Pt(9.0)
                r.font.color.rgb = RGBColor(0x0F, 0x2A, 0x44)

            sp = pPr.find(qn("w:spacing"))
            if sp is not None:
                sp.set(qn("w:before"), "20")
                sp.set(qn("w:after"), "20")
                sp.set(qn("w:line"), "252")
                sp.set(qn("w:lineRule"), "auto")
            else:
                pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="20" w:after="20" w:line="252" w:lineRule="auto"/>'))

            if pPr.find(qn("w:shd")) is None:
                pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
            if pPr.find(qn("w:pBdr")) is None:
                bdr = parse_xml(f'''<w:pBdr {nsdecls("w")}>
                    <w:left w:val="single" w:sz="6" w:space="8" w:color="CBD5E1"/>
                    <w:top w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                    <w:bottom w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                    <w:right w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                </w:pBdr>''')
                pPr.append(bdr)

            para.paragraph_format.left_indent = Inches(0.15)
            para.paragraph_format.right_indent = Inches(0.15)

        # Blockquote
        elif style_name == "Block Text" or (para._p.pPr is not None and para._p.pPr.find(qn("w:pStyle")) is not None and para._p.pPr.find(qn("w:pStyle")).attrib.get(qn("w:val")) == "BlockText"):
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pPr = para._p.get_or_add_pPr()
            if pPr.find(qn("w:shd")) is None:
                pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'))
            if pPr.find(qn("w:pBdr")) is None:
                bdr = parse_xml(f'''<w:pBdr {nsdecls("w")}>
                    <w:left w:val="single" w:sz="12" w:space="8" w:color="000000"/>
                    <w:top w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                    <w:bottom w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                    <w:right w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>
                </w:pBdr>''')
                pPr.append(bdr)
            para.paragraph_format.left_indent = Inches(0.2)
            para.paragraph_format.right_indent = Inches(0.2)
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(4)
            for r in para.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.5)
                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Body text
        else:
            if idx in range(1, 4):  # Lời nói đầu paragraphs
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.line_spacing = 1.5
                para.paragraph_format.space_after = Pt(6)
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(14)
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.line_spacing = 1.15
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(3)
                for r in para.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
                    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 4. Gắn hình ảnh minh họa bối cảnh vào từng bài toán
    print(f"🖼️ Đang chèn hình ảnh minh họa trực quan bối cảnh cho Quyển 3...")
    tap3_codes = set(p["code"] for l in manifest_data["lessons"][14:21] for p in l["problems"])
    for prob_code in tap3_codes:
        if prob_code not in DIAGRAMS_MAP:
            continue
        img_filename, caption_text = DIAGRAMS_MAP[prob_code]
        h3_idx = None
        for i, p in enumerate(doc.paragraphs):
            if prob_code in p.text and p.style.name == "Heading 3":
                h3_idx = i
                break
        if h3_idx is None:
            continue

        target_p = None
        for j in range(h3_idx + 1, min(len(doc.paragraphs), h3_idx + 6)):
            p_curr = doc.paragraphs[j]
            if p_curr.style.name == "Heading 3":
                break
            if p_curr.text.strip().startswith("Bối cảnh:"):
                next_p = doc.paragraphs[j + 1] if j + 1 < len(doc.paragraphs) else p_curr
                target_p = next_p
                break
        if target_p is None:
            target_p = doc.paragraphs[h3_idx + 1] if h3_idx + 1 < len(doc.paragraphs) else doc.paragraphs[h3_idx]

        insert_picture_with_caption(doc, img_filename, caption_text, target_p._p)

    # 5. Tables Processing & Dynamic Left Indent & Spacing 120 dxa
    body_element = doc._body._element
    for i, child in enumerate(body_element):
        if child.tag.split("}")[-1] == "tbl":
            prev_el = body_element[i - 1] if i > 0 else None
            if prev_el is not None and prev_el.tag.split("}")[-1] == "p":
                prev_pPr = prev_el.find(qn("w:pPr"))
                if prev_pPr is not None:
                    sp = prev_pPr.find(qn("w:spacing"))
                    if sp is not None:
                        curr_after = sp.attrib.get(qn("w:after"), None)
                        if curr_after is None or int(curr_after) < 120:
                            sp.set(qn("w:after"), "120")
                    else:
                        prev_pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:after="120"/>'))

            next_el = body_element[i + 1] if i + 1 < len(body_element) else None
            if next_el is not None and next_el.tag.split("}")[-1] == "p":
                next_pPr = next_el.find(qn("w:pPr"))
                if next_pPr is not None:
                    sp = next_pPr.find(qn("w:spacing"))
                    if sp is not None:
                        curr_before = sp.attrib.get(qn("w:before"), None)
                        if curr_before is None or int(curr_before) < 120:
                            sp.set(qn("w:before"), "120")
                    else:
                        next_pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="120"/>'))

    for table in doc.tables:
        table_tblPr = table._tbl.tblPr

        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            for h in trPr.findall(qn("w:tblHeader")): trPr.remove(h)
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        for mr in table._tbl.findall(".//{" + math_ns + "}r"):
            for m_rPr in mr.findall("{" + math_ns + "}rPr"):
                mr.remove(m_rPr)
            mrPr = parse_xml(f'<m:rPr {nsdecls("m")}><m:scr m:val="roman"/><m:sty m:val="p"/><w:rPr {nsdecls("w")}><w:sz w:val="24"/></w:rPr></m:rPr>')
            mr.append(mrPr)

        first_row_text = "".join(c.text for c in table.rows[0].cells).strip().lower()
        is_sample_table = "đầu vào" in first_row_text and "đầu ra" in first_row_text

        if is_sample_table:
            for w_elem in table_tblPr.findall(qn("w:tblW")): table_tblPr.remove(w_elem)
            table_tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="6800" w:type="dxa"/>'))
            for jc in table_tblPr.findall(qn("w:jc")): table_tblPr.remove(jc)
            table_tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))

            for b in table_tblPr.findall(qn("w:tblBorders")): table_tblPr.remove(b)
            table_tblPr.append(parse_xml(f'''<w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            </w:tblBorders>'''))

            # Row 0: Header
            row0 = table.rows[0]
            for cell in row0.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                for w in tcPr.findall(qn("w:tcW")): tcPr.remove(w)
                tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="3400" w:type="dxa"/>'))
                for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
                for tm in tcPr.findall(qn("w:tcMar")): tcPr.remove(tm)
                tcPr.append(parse_xml(f'''<w:tcMar {nsdecls("w")}>
                    <w:top w:w="60" w:type="dxa"/>
                    <w:left w:w="80" w:type="dxa"/>
                    <w:bottom w:w="60" w:type="dxa"/>
                    <w:right w:w="80" w:type="dxa"/>
                </w:tcMar>'''))
                for va in tcPr.findall(qn("w:vAlign")): tcPr.remove(va)
                tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

                for p in cell.paragraphs:
                    pPr = p._p.get_or_add_pPr()
                    sp = pPr.find(qn("w:spacing"))
                    if sp is not None:
                        sp.set(qn("w:before"), "60")
                        sp.set(qn("w:after"), "60")
                        sp.set(qn("w:line"), "276")
                        sp.set(qn("w:lineRule"), "auto")
                    else:
                        pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60" w:line="276" w:lineRule="auto"/>'))
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Consolas"
                        r.font.size = Pt(11)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

            # Row 1+: Testcase data
            for r_idx in range(1, len(table.rows)):
                row = table.rows[r_idx]
                for cell in row.cells:
                    tcPr = cell._tc.get_or_add_tcPr()
                    for w in tcPr.findall(qn("w:tcW")): tcPr.remove(w)
                    tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="3400" w:type="dxa"/>'))
                    for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>'))
                    for tm in tcPr.findall(qn("w:tcMar")): tcPr.remove(tm)
                    tcPr.append(parse_xml(f'''<w:tcMar {nsdecls("w")}>
                        <w:top w:w="40" w:type="dxa"/>
                        <w:left w:w="60" w:type="dxa"/>
                        <w:bottom w:w="40" w:type="dxa"/>
                        <w:right w:w="60" w:type="dxa"/>
                    </w:tcMar>'''))
                    for va in tcPr.findall(qn("w:vAlign")): tcPr.remove(va)
                    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

                    orig_p_list = list(cell.paragraphs)
                    lines = []
                    for p in orig_p_list:
                        p_text = p.text
                        if "<br>" in p_text or "<br/>" in p_text or "&lt;br&gt;" in p_text:
                            sub = re.split(r'<br\s*/?>|&lt;br&gt;', p_text)
                            for s_str in sub:
                                s_clean = s_str.strip()
                                if s_clean: lines.append(s_clean)
                        else:
                            p_clean = p_text.strip()
                            if p_clean: lines.append(p_clean)

                    if not lines: lines = [""]

                    max_len = max(len(l) for l in lines) if lines else 0
                    if max_len <= 4:
                        dyn_indent_pt = 72.0
                    else:
                        dyn_indent_pt = max(10.0, 72.0 - (max_len - 4) * 3.25)
                    dyn_indent_dxa = int(round(dyn_indent_pt * 20))

                    for p in orig_p_list:
                        p_elem = p._p
                        if p_elem.getparent() is not None:
                            p_elem.getparent().remove(p_elem)

                    for line_str in lines:
                        new_p = parse_xml(f'''<w:p {nsdecls("w")}>
                            <w:pPr>
                                <w:jc w:val="left"/>
                                <w:ind w:left="{dyn_indent_dxa}"/>
                                <w:spacing w:before="30" w:after="30" w:line="240" w:lineRule="auto"/>
                            </w:pPr>
                            <w:r>
                                <w:rPr>
                                    <w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>
                                    <w:sz w:val="19"/>
                                    <w:szCs w:val="19"/>
                                    <w:color w:val="0F2A44"/>
                                </w:rPr>
                                <w:t xml:space="preserve">{html.escape(line_str)}</w:t>
                            </w:r>
                        </w:p>''')
                        cell._tc.append(new_p)

        else:
            # Dry Run Table
            table_tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
            for b in table_tblPr.findall(qn("w:tblBorders")): table_tblPr.remove(b)
            table_tblPr.append(parse_xml(f'''<w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            </w:tblBorders>'''))

            for r_idx, row in enumerate(table.rows):
                for cell in row.cells:
                    tcPr = cell._tc.get_or_add_tcPr()
                    for tm in tcPr.findall(qn("w:tcMar")): tcPr.remove(tm)
                    tcPr.append(parse_xml(f'''<w:tcMar {nsdecls("w")}>
                        <w:top w:w="60" w:type="dxa"/>
                        <w:left w:w="80" w:type="dxa"/>
                        <w:bottom w:w="60" w:type="dxa"/>
                        <w:right w:w="80" w:type="dxa"/>
                    </w:tcMar>'''))
                    for va in tcPr.findall(qn("w:vAlign")): tcPr.remove(va)
                    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

                    if r_idx == 0:
                        for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
                        for p in cell.paragraphs:
                            pPr = p._p.get_or_add_pPr()
                            sp = pPr.find(qn("w:spacing"))
                            if sp is not None:
                                sp.set(qn("w:before"), "60")
                                sp.set(qn("w:after"), "60")
                                sp.set(qn("w:line"), "276")
                                sp.set(qn("w:lineRule"), "auto")
                            else:
                                pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60" w:line="276" w:lineRule="auto"/>'))
                            for r in p.runs:
                                r.font.name = "Times New Roman"
                                r.font.size = Pt(12)
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
                    else:
                        fill_c = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
                        for s_elem in tcPr.findall(qn("w:shd")): tcPr.remove(s_elem)
                        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_c}"/>'))
                        for p in cell.paragraphs:
                            pPr = p._p.get_or_add_pPr()
                            sp = pPr.find(qn("w:spacing"))
                            if sp is not None:
                                sp.set(qn("w:before"), "40")
                                sp.set(qn("w:after"), "40")
                                sp.set(qn("w:line"), "276")
                                sp.set(qn("w:lineRule"), "auto")
                            else:
                                pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="40" w:after="40" w:line="276" w:lineRule="auto"/>'))
                            for r in p.runs:
                                r.font.name = "Times New Roman"
                                r.font.size = Pt(12)
                                r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 6. Build Interactive Hyperlinked TOC at End (Chương + Bài học)
    if p_toc_heading:
        print(f"  → Tạo Mục lục tương tác gồm {len(toc_headings)} liên kết...")
        last_elem = p_toc_heading._p

        for item in toc_headings:
            title = html.escape(item["text"])
            bm = item["bm_name"]

            if item.get("is_chapter"):
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:spacing w:before="180" w:after="40"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:b/>
                                <w:sz w:val="23"/>
                                <w:color w:val="0F2A44"/>
                            </w:rPr>
                            <w:t>{title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            elif item.get("is_lesson"):
                p_xml = f'''<w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:ind w:left="280"/>
                        <w:spacing w:before="20" w:after="20"/>
                    </w:pPr>
                    <w:hyperlink w:anchor="{bm}">
                        <w:r>
                            <w:rPr>
                                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                                <w:sz w:val="22"/>
                                <w:color w:val="1A4A6B"/>
                            </w:rPr>
                            <w:t>•  {title}</w:t>
                        </w:r>
                    </w:hyperlink>
                </w:p>'''
            else:
                continue

            p_elem = parse_xml(p_xml)
            last_elem.addnext(p_elem)
            last_elem = p_elem

    # 7. Lưu file docx hoàn thiện
    doc.save(str(docx_path))
    print(f"  ✅ Đã lưu file: {docx_path.name}")

    # 8. Hậu kiểm dọn sạch package XML & ép màu đen 000000 trong styles.xml
    clean_and_blacken_package(docx_path)

    # 9. Đồng bộ Header & Đường kẻ ngang Straight Connector 1 chuẩn OpenXML
    sync_headers(docx_path, "Giáo trình C++ - Sách giáo viên - Quyển 3")

def clean_and_blacken_package(docx_path):
    print("🧹 Đang dọn dẹp package XML, ép đen 000000 và chuẩn hóa docPr ID...")
    temp_zip_path = docx_path.parent / f"_temp_clean_{docx_path.name}"
    color_pat = re.compile(r'(<w:color\b[^>]*?\bw:val=[\'"])1E293B([\'"][^>]*?>)', re.IGNORECASE)

    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(temp_zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/styles.xml":
                    text = data.decode("utf-8")
                    text, _ = color_pat.subn(r'\g<1>000000\g<2>', text)
                    text = re.sub(r'w:color\s+w:themeColor="[^"]*"\s+w:themeShade="[^"]*"\s+w:val="000000"', 'w:color w:val="000000"', text)
                    data = text.encode("utf-8")
                elif item.filename == "word/document.xml":
                    root = ET.fromstring(data)
                    curr_id = 1000
                    for d in root.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
                        d.set("id", str(curr_id))
                        curr_id += 1
                    data = ET.tostring(root, xml_declaration=True, encoding="utf-8")
                    text = data.decode("utf-8")
                    text, _ = color_pat.subn(r'\g<1>000000\g<2>', text)
                    data = text.encode("utf-8")
                zout.writestr(item, data)

    temp_zip_path.replace(docx_path)
    print(f"  ✅ Hoàn tất ép đen, chuẩn hóa docPr ID và dọn sạch XML cho {docx_path.name}")

def sync_headers(docx_path, header_title):
    print(f"📏 Đang đồng bộ Header và Đường kẻ ngang Straight Connector 1: {header_title}...")
    temp_zip = docx_path.parent / f"_temp_sync_hdr_{docx_path.name}"
    
    # Lấy mẫu header3 chuẩn từ cpp-giaovien-quyen-1.docx
    ref_docx = BASE_DIR / "cpp-giaovien-quyen-1.docx"
    with zipfile.ZipFile(ref_docx, "r") as z_ref:
        raw_h3 = z_ref.read("word/header3.xml")

    root_h3 = ET.fromstring(raw_h3)
    p0 = root_h3.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
    # Loại bỏ tab thừa nếu có, chỉ giữ đúng 1 tab căn phải chuẩn pos="9899"
    tab_runs = [r for r in p0.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r") if r.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tab") is not None]
    if len(tab_runs) > 1:
        for extra_tr in tab_runs[1:]:
            p0.remove(extra_tr)

    # Đổi text sang tiêu đề tập tương ứng
    for t in p0.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
        if t.text and "Giáo trình C++" in t.text:
            t.text = header_title

    h3_xml_clean = ET.tostring(root_h3, xml_declaration=True, encoding="utf-8").decode("utf-8")

    # Header 1 (trang chẵn)
    h1_xml = h3_xml_clean.replace("WordPictureWatermark1286990221", "WordPictureWatermark1286990222")
    h1_xml = h1_xml.replace("_x0000_s2049", "_x0000_s2052")
    h1_xml = h1_xml.replace('id="318338836"', 'id="318338837"')

    # Header 2 (trang lẻ / default)
    h2_xml = h3_xml_clean.replace("WordPictureWatermark1286990221", "WordPictureWatermark1286990223")
    h2_xml = h2_xml.replace("_x0000_s2049", "_x0000_s2051")
    h2_xml = h2_xml.replace('id="318338836"', 'id="318338838"')

    # Header 3 (trang đầu)
    h3_xml = h3_xml_clean

    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(temp_zip, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/header1.xml":
                    zout.writestr(item, h1_xml.encode("utf-8"))
                elif item.filename == "word/header2.xml":
                    zout.writestr(item, h2_xml.encode("utf-8"))
                elif item.filename == "word/header3.xml":
                    zout.writestr(item, h3_xml.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))

    temp_zip.replace(docx_path)
    print("  ✅ Đã đồng bộ sạch toàn bộ 3 Headers có đường kẻ ngang chuẩn OpenXML!")

def main():
    print("================================================================================")
    print("🚀 BẮT ĐẦU XUẤT BẢN SÁCH GIÁO VIÊN C++ TẬP 3 (BÀI 15 - 21, 105 BÀI TOÁN)")
    print("================================================================================")

    with open(MANIFEST_FILE, "r", encoding="utf-8") as f:
        manifest_data = json.load(f)

    # 1. Sinh Markdown
    full_md = build_tap3_markdown(manifest_data)

    # 2. Convert to DOCX qua Pandoc
    convert_markdown_to_docx(full_md, OUT_DOCX)

    # 3. Post-process Design System (Ảnh minh họa, bảng, màu sắc, TOC, Header)
    post_process_teacher_docx(OUT_DOCX, manifest_data)

    print("\n================================================================================")
    print(f"🎉 XUẤT BẢN THÀNH CÔNG: {OUT_DOCX.name} ({OUT_DOCX.stat().st_size:,} bytes)!")
    print("================================================================================")

if __name__ == "__main__":
    main()
